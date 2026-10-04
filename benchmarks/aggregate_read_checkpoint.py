"""Joint managed policies on equivalent shared text/formula/styled-date values."""
import argparse
from datetime import datetime
import io
import json
import os
from pathlib import Path
import platform
import statistics
import subprocess
import sys
import tempfile
import zipfile
import openpyxl
from shared_strings_checkpoint import ROOT, HERE, run
from iso_checkpoint import measure

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
TAIL = "X" * 96

def generate(path, rows):
    book = openpyxl.Workbook()
    book.active["C1"].number_format = "yyyy-mm-dd"
    buffer = io.BytesIO()
    book.save(buffer)
    with zipfile.ZipFile(buffer) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(b"</Relationships>", b'<Relationship Id="shared" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(b"</Types>", b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            if name != "xl/worksheets/sheet1.xml":
                archive.writestr(name, value)
        with archive.open("xl/sharedStrings.xml", "w") as source:
            source.write(f'<sst xmlns="{MAIN}">'.encode())
            for index in range(1, rows + 1):
                source.write(f'<si><t>text-{index:06}-{TAIL}</t></si>'.encode())
            source.write(b'</sst>')
        with archive.open("xl/worksheets/sheet1.xml", "w") as source:
            source.write(f'<worksheet xmlns="{MAIN}"><sheetData>'.encode())
            for index in range(1, rows + 1):
                group = (index - 1) // 16
                first = group * 16 + 1
                formula = f'<f t="shared" si="{group}" ref="B{first}:B{min(first+15, rows)}">B{first}+1</f>' if index == first else f'<f t="shared" si="{group}"/>'
                source.write(f'<row r="{index}"><c r="A{index}" t="s"><v>{index-1}</v></c><c r="B{index}">{formula}<v>{index}</v></c><c r="C{index}" s="1"><v>43831</v></c></row>'.encode())
            source.write(b'</sheetData></worksheet>')

def reference(path, mode):
    book = openpyxl.load_workbook(path, read_only=mode == "scan")
    count = 0
    for index, row in enumerate(book.active, 1):
        assert [cell.value for cell in row] == [f"text-{index:06}-{TAIL}", f"=B{index}+1", datetime(2020, 1, 1)]
        count += 3
    book.close()
    print(count)

def stats(command):
    result = run(command)
    return json.loads(next(line.removeprefix("POLICY_STATS ") for line in result.stderr.splitlines() if line.startswith("POLICY_STATS ")))

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows", type=int, nargs="+", default=[10000, 100000])
    parser.add_argument("--runs", type=int, default=5)
    parser.add_argument("--baseline", type=Path, required=True)
    parser.add_argument("--baseline-core", required=True)
    parser.add_argument("--output", type=Path, default=HERE / "results/m2-aggregate-read.json")
    args = parser.parse_args()
    assert openpyxl.__version__ == "3.1.5"
    if args.runs < 1 or any(not 0 < size <= 1048576 for size in args.rows):
        parser.error("Require positive samples and physical worksheet rows")
    target = Path(os.environ.get("CARGO_TARGET_DIR", ROOT / "target"))
    report = {"reference": openpyxl.__version__, "baseline_core": args.baseline_core, "platform": platform.platform(), "measurement": f"One warmup and {args.runs} rotating serial wall/CPU/RSS samples; no builds/checks during timing", "semantics": "All readers verify every shared text, expanded formula and styled date. Native additionally verifies cached integers in the same pass; public cached values use a separate untimed pass. Native narrow/wide policies receive 2/256MiB budgets; old budget governs rows only while current jointly bounds managed catalogs/cache/templates/library-retained rows. Scan reference is read_only; repeated-wide reference is a full model; repeated-tight reference is read_only to match final fallback. Native calendar serial/epoch/kind correspond to reference datetime, with prior public date mapping verification. Default theme is preloaded in both native runs. Inspection/invariants are checked outside timing. Exact anonymous string-spool storage is obtained from native stats because directory sampling misses unlinked files; OS page cache, allocator/dependency overhead and caller-owned output are additional. No calamine/rust_xlsxwriter overlap timing is claimed here.", "cases": []}
    (HERE / "data").mkdir(exist_ok=True)
    for size in args.rows:
        path = HERE / "data" / f"aggregate-mixed-{size}.xlsx"
        generate(path, size)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        checked = 0
        for index, row in enumerate(cached.active, 1):
            assert [cell.value for cell in row] == [f"text-{index:06}-{TAIL}", index, datetime(2020, 1, 1)]
            checked += 3
        cached.close()
        assert checked == size * 3
        for mode, budget in [("scan", 2), ("scan", 256), ("repeated", 2), ("repeated", 256)]:
            commands = {"crabxl": [target / "release/examples/adaptive_mixed_read", path, mode, budget], "crabxl-before": [args.baseline.resolve(), path, mode, budget], "openpyxl": [sys.executable, __file__, "--reference", path, "repeated" if mode == "repeated" and budget == 256 else "scan"]}
            inspected = stats(commands["crabxl"] + ["inspect"])
            previous = stats(commands["crabxl-before"])
            expected_mode = "Materialized" if mode == "repeated" and budget == 256 else "Streaming"
            assert inspected["mode"] == previous["mode"] == expected_mode
            assert inspected["inspected_managed_bytes"] <= inspected["budget"] - inspected["working"]
            assert inspected["disk_backed"] == (budget == 2)
            samples = {name: [] for name in commands}
            with tempfile.TemporaryDirectory() as directory:
                for command in commands.values():
                    output, _ = measure(command, Path(directory))
                    assert output == str(checked)
                names = list(commands)
                for iteration in range(args.runs):
                    for name in names[iteration % len(names):] + names[:iteration % len(names)]:
                        output, sample = measure(commands[name], Path(directory))
                        assert output == str(checked)
                        samples[name].append(sample)
                        print(size, mode, budget, name, sample, flush=True)
            report["cases"].append({"rows": size, "cells": checked, "access": mode, "budget_mib": budget, "input_bytes": path.stat().st_size, "current_policy_stats": inspected, "prior_policy_stats": previous, "samples": samples, "medians": {name: {key: statistics.median(sample[key] for sample in values) for key in ("seconds", "cpu_seconds", "peak_rss_kib", "sampled_temp_peak_bytes")} for name, values in samples.items()}})
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(json.dumps(report, indent=2) + "\n")

if __name__ == "__main__":
    if len(sys.argv) == 4 and sys.argv[1] == "--reference":
        reference(sys.argv[2], sys.argv[3])
    else:
        main()
