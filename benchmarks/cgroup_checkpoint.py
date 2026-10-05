"""Serial numeric Auto-policy regression after mounted-controller discovery."""
import argparse
import hashlib
import json
from pathlib import Path
import statistics
import sys
import tempfile
import openpyxl
from iso_checkpoint import measure


def reference(path, mode):
    book = openpyxl.load_workbook(path, read_only=mode == 'auto-scan', data_only=True)
    count = total = 0
    for _ in range(1 if mode == 'auto-scan' else 3):
        for values in book.active.iter_rows(values_only=True):
            for value in values:
                count += 1
                total += value
    book.close()
    print(count, total)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--baseline', type=Path, required=True)
    parser.add_argument('--current', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=Path(__file__).parent / 'results/m2-cgroup-discovery.json')
    args = parser.parse_args()
    assert openpyxl.__version__ == '3.1.5'
    report = {'baseline_core': 'fa5f3b0fb1148803ee312de726fd182be152e504', 'baseline_binary_sha256': hashlib.sha256(args.baseline.read_bytes()).hexdigest(), 'reference': openpyxl.__version__, 'measurement': 'One warmup and five rotating serial cold-process runs; builds/tests/generation excluded. Kernel RSS, wall/CPU; 25ms temporary sampling with cleanup. Auto decisions retained per sample.', 'semantics': 'Ten numeric columns. Auto scan streams once; Auto repeated access retains a sparse snapshot when it fits and sums three times. Corresponding reference uses read_only scanning or ordinary loaded model. Full row count and checksum checked. Native processes and Python runtime overhead differ. Live host is cgroup v2; v1 mount/hierarchy behavior is deterministic fixture-tested, not performance-tested on an actual v1 host. Availability is a snapshot and may vary per process.', 'cases': []}
    with tempfile.TemporaryDirectory() as name:
        root = Path(name)
        temporary = root / 'temporary'
        temporary.mkdir()
        for rows in (5000, 50000):
            path = root / 'numbers.xlsx'
            book = openpyxl.Workbook(write_only=True)
            sheet = book.create_sheet('Sheet')
            for row in range(rows):
                sheet.append(list(range(row * 10, row * 10 + 10)))
            book.save(path)
            book.close()
            cells = rows * 10
            for mode in ('auto-scan', 'auto-repeat'):
                passes = 1 if mode == 'auto-scan' else 3
                expected = f'{cells * passes} {cells * (cells - 1) // 2 * passes}'
                commands = {engine: [binary.resolve(), path, 'Sheet', 32768, mode, 'auto', passes] for engine, binary in [('current', args.current), ('baseline', args.baseline)]}
                commands['openpyxl'] = [sys.executable, __file__, '--reference', path, mode]
                samples = {engine: [] for engine in commands}
                for command in commands.values():
                    result, _ = measure(command, temporary)
                    assert result == expected
                engines = list(commands)
                for iteration in range(5):
                    for engine in engines[iteration % 3:] + engines[:iteration % 3]:
                        result, sample = measure(commands[engine], temporary)
                        assert result == expected
                        samples[engine].append(sample)
                report['cases'].append({'rows': rows, 'cells': cells, 'mode': mode, 'passes': passes, 'expected': expected, 'source_bytes': path.stat().st_size, 'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'samples': samples, 'medians': {engine: {key: statistics.median(item[key] for item in values) for key in ('seconds', 'cpu_seconds', 'peak_rss_kib', 'sampled_temp_peak_bytes')} for engine, values in samples.items()}})
                args.output.write_text(json.dumps(report, indent=2) + '\n')


if __name__ == '__main__':
    if len(sys.argv) > 1 and sys.argv[1] == '--reference':
        reference(sys.argv[2], sys.argv[3])
    else:
        main()
