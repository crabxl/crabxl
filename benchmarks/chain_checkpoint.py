"""Derived calculation-chain removal, edit interoperability and regression evidence."""
import argparse
import io
import xml.etree.ElementTree as ET
import json
import os
from pathlib import Path
import re
import statistics
import subprocess
import sys
import tempfile
import time
import zipfile
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--before-editor',type=Path,required=True)
args=parser.parse_args()
report={'scope':'Default derived calculation-chain discard, explicit rejection policy, coordinated part/relationship/content-type removal; not loaded structural editing','versions':'Rust 1.88 release, openpyxl 3.1.5, Python 3.12; before editor at 1f7e870/d576595 equivalent edit path','cases':[]}
def measure(command,directory):
    p=subprocess.Popen([str(ROOT/'benchmarks/measure'),*map(str,command)],env=dict(os.environ,TMPDIR=str(directory)),stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True)
    peak=0
    while p.poll() is None:
        total=0
        for f in directory.iterdir():
            if f.name.endswith('.xlsx'):continue
            try:total+=f.stat().st_size
            except FileNotFoundError:pass
        peak=max(peak,total);time.sleep(.01)
    out,err=p.communicate()
    if p.returncode:raise RuntimeError(err)
    result=json.loads(err.split('MEASURE ')[-1]);result.update(stdout=out.strip(),observed_extra_temp_bytes=peak)
    return result
for rows in (10000,100000):
    source=ROOT/'benchmarks/data'/f'numbers-{rows}.xlsx'
    if not source.exists():
        book=openpyxl.Workbook(write_only=True);sheet=book.create_sheet('Sheet')
        for row in range(rows):sheet.append([row*10+c for c in range(10)])
        book.save(source)
    with tempfile.TemporaryDirectory(prefix='crabxl-chain-') as name:
        directory=Path(name);chained=directory/'source.xlsx';target=directory/'output.xlsx'
        with zipfile.ZipFile(source) as original,zipfile.ZipFile(chained,'w',zipfile.ZIP_DEFLATED) as out:
            replacements=0
            def formula(match):
                global replacements
                replacements+=1
                row,value=match.groups()
                return b'<c r="B'+row+b'"><f>A'+row+b'+1</f><v>'+value+b'</v></c>'
            for entry in original.infolist():
                data=original.read(entry.filename)
                if entry.filename=='xl/worksheets/sheet1.xml':data=re.sub(rb'<c r="B(\d+)"[^>]*><v>(\d+)</v></c>',formula,data)
                if entry.filename=='[Content_Types].xml':data=data.replace(b'</Types>',b'<Override PartName="/custom/order.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>')
                if entry.filename=='xl/_rels/workbook.xml.rels':data=data.replace(b'</Relationships>',b'<Relationship Id="chain" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="../custom/order.xml"/></Relationships>')
                out.writestr(entry.filename,data)
            chain=b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'+b''.join(f'<c r="B{row+1}" i="1"/>'.encode() for row in range(rows))+b'</calcChain>'
            out.writestr('custom/order.xml',chain)
        assert replacements==rows
        for workload,commands in [
            ('chain-edit',{'crabxl-native':[ROOT/'target/release/examples/edit_demo',chained,target,'edit'],'openpyxl':[sys.executable,ROOT/'benchmarks/edit_baseline.py',chained,target]}),
            ('editor-regression',{'before':[args.before_editor,source,target,'edit'],'current':[ROOT/'target/release/examples/edit_demo',source,target,'edit']}),
        ]:
            for cmd in commands.values():measure(cmd,directory)
            runs=[]
            for iteration in range(3):
                for engine in (list(commands) if iteration%2==0 else list(reversed(commands))):
                    run=measure(commands[engine],directory)
                    total=rows*10;expected=total*(total-1)//2-1
                    if workload=='chain-edit':expected-=5*rows*(rows-1)+rows
                    count=total-rows if workload=='chain-edit' else total
                    checksum=subprocess.run([ROOT/'target/release/examples/sum',target],check=True,capture_output=True,text=True).stdout.strip()
                    assert checksum==f'{count} {expected}'
                    with zipfile.ZipFile(target) as archive:
                        assert 'custom/order.xml' not in archive.namelist()
                        assert b'calcChain' not in archive.read('[Content_Types].xml')
                        assert b'calcChain' not in archive.read('xl/_rels/workbook.xml.rels')
                        if workload=='chain-edit':
                            data=archive.read('xl/worksheets/sheet1.xml');assert data.count(b'<f>')==rows
                            for _, node in ET.iterparse(io.BytesIO(data), events=('end',)):
                                if node.tag=='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c':
                                    formula=node.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f')
                                    cache=node.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v')
                                    if formula is not None: assert cache is None or not cache.text
                                    node.clear()
                    if workload=='chain-edit':
                        book=openpyxl.load_workbook(target,read_only=True)
                        seen=0
                        for row in book['Sheet'].values:
                            expected_row=[seen*10+c for c in range(10)];expected_row[1]=f'=A{seen+1}+1'
                            if seen==0:expected_row[0]=-1
                            assert row==tuple(expected_row);seen+=1
                        assert seen==rows;book.close()
                    run.update(engine=engine,checksum=checksum,output_bytes=target.stat().st_size,cleanup=all(p.suffix=='.xlsx' for p in directory.iterdir()));assert run['cleanup'];runs.append(run)
            report['cases'].append({'workload':workload,'rows':rows,'columns':10,'chain_xml_bytes':len(chain),'runs':runs})
            print(rows,workload,[(engine,statistics.median(r['seconds'] for r in runs if r['engine']==engine)) for engine in commands],flush=True)
(ROOT/'benchmarks/results/m4-chain.json').write_text(json.dumps(report,indent=2)+'\n')
