"""Equivalent projected dynamic array workloads; native also checks every cache."""
import argparse
import json
import os
from pathlib import Path
import platform
import statistics
import sys
import tempfile
import zipfile
import openpyxl
from shared_strings_checkpoint import ROOT, HERE, run
from openpyxl.worksheet.formula import ArrayFormula
from iso_checkpoint import measure


def generate(path, rows):
    book = openpyxl.Workbook()
    book.save(path)
    book.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            if name != "xl/worksheets/sheet1.xml":
                archive.writestr(name, value)
        with archive.open("xl/worksheets/sheet1.xml", "w") as sheet:
            sheet.write(b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>')
            for index in range(1, rows + 1):
                sheet.write(f'<row r="{index}"><c r="A{index}" cm="1"><f t="array" ref="A{index}">_xlfn.SEQUENCE(1)</f><v>{index}</v></c><c r="B{index}" vm="1"><v>{index}</v></c></row>'.encode())
            sheet.write(b'</sheetData></worksheet>')


def reference(path):
    book = openpyxl.load_workbook(path, read_only=True)
    count = 0
    for index, row in enumerate(book.active, 1):
        assert len(row) == 2
        assert isinstance(row[0].value, ArrayFormula)
        assert row[0].value.text == '=_xlfn.SEQUENCE(1)'
        assert row[0].value.ref == f'A{index}'
        assert row[1].value == index
        count += 2
    book.close()
    print(count)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows", type=int, nargs="+", default=[10000, 100000])
    parser.add_argument("--runs", type=int, default=5)
    parser.add_argument("--output", type=Path, default=HERE / "results/m2-dynamic-formulas.json")
    args = parser.parse_args()
    assert openpyxl.__version__ == "3.1.5"
    if args.runs < 1 or any(not 0 < size <= 1048576 for size in args.rows):
        parser.error("Require positive samples and physical worksheet row bounds")
    run(["cc", "-O2", "-Wall", "-Wextra", "-Werror", HERE / "measure.c", "-o", HERE / "measure"])
    run(["cargo", "build", "--release", "--locked", "-p", "crabxl", "--example", "dynamic_formula_read"])
    target = Path(os.environ.get("CARGO_TARGET_DIR", ROOT / "target"))
    report = {"reference": openpyxl.__version__, "platform": platform.platform(), "measurement": f"One warmup and {args.runs} rotating serial wall/CPU/RSS samples including process baseline", "semantics": "Both readers stream and verify every ArrayFormula expression/range and scalar. Native also checks every array cache in the same pass; public cache validation is a separate untimed data_only pass. Opaque cm/vm graph references are projected like the public baseline, not interpreted. No temporary storage or worksheet materialization. Prior native revision rejects these inputs; no prior dynamic-input timing is possible. calamine formula/cache-only overlap needs separate verification and is not claimed here.", "cases": []}
    (HERE / "data").mkdir(exist_ok=True)
    for size in args.rows:
        path = HERE / "data" / f"dynamic-formulas-{size}.xlsx"
        generate(path, size)
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        checked = 0
        for index, row in enumerate(book.active, 1):
            assert [cell.value for cell in row] == [index, index]
            checked += 2
        book.close()
        assert checked == size * 2
        commands = {"crabxl": [target / "release/examples/dynamic_formula_read", path], "openpyxl": [sys.executable, __file__, "--reference", path]}
        samples = {name: [] for name in commands}
        with tempfile.TemporaryDirectory() as directory:
            for command in commands.values():
                output, _ = measure(command, Path(directory))
                assert output == str(checked)
            names = list(commands)
            for iteration in range(args.runs):
                for name in names[iteration % len(names):] + names[:iteration % len(names)]:
                    output, sample = measure(commands[name], Path(directory))
                    assert output == str(checked)
                    assert sample["sampled_temp_peak_bytes"] == 0
                    samples[name].append(sample)
                    print(size, name, sample, flush=True)
        report["cases"].append({"rows": size, "cells": checked, "input_bytes": path.stat().st_size, "samples": samples, "medians": {name: {key: statistics.median(s[key] for s in values) for key in ("seconds", "cpu_seconds", "peak_rss_kib", "sampled_temp_peak_bytes")} for name, values in samples.items()}})
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(report, indent=2) + "\n")


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--reference":
        reference(sys.argv[2])
    else:
        main()
