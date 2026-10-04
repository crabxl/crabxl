"""M4 lazy preservation/editing, sparse-model costs and M3 reader regression."""
import argparse
import datetime
import io
import json
import os
from pathlib import Path
import platform
import statistics
import subprocess
import sys
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.styles import Font
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image
from PIL import Image as PillowImage

ROOT=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--before-reader',required=True,help='Release M3 sum executable at 4aaaee7')
args=parser.parse_args()
assert openpyxl.__version__=='3.1.5'
(ROOT/'benchmarks/data').mkdir(exist_ok=True)

def measure(command,directory=None):
    environment=dict(os.environ)
    if directory: environment['TMPDIR']=str(directory)
    process=subprocess.Popen([str(ROOT/'benchmarks/measure'),*map(str,command)],env=environment,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True)
    temporary=0
    while process.poll() is None:
        if directory:
            total=0
            for path in directory.iterdir():
                if path.suffix=='.xlsx': continue
                try: total+=path.stat().st_size
                except FileNotFoundError: pass
            temporary=max(temporary,total)
        time.sleep(.025)
    stdout,stderr=process.communicate()
    if process.returncode: raise RuntimeError(stderr)
    result={'stdout':stdout.strip(),**json.loads(stderr.split('MEASURE ')[-1])}
    if directory:
        result['observed_temp_bytes']=temporary
        assert all(path.suffix=='.xlsx' for path in directory.iterdir()),'Output/worksheet temporary files leaked'
    return result

def checksum(path,rows,edited):
    cells=rows*10
    value=subprocess.run([str(ROOT/'target/release/examples/sum'),str(path)],check=True,capture_output=True,text=True).stdout.strip()
    assert value==f'{cells} {cells*(cells-1)//2-int(edited)}',value

def public_readback(path,rows):
    book=openpyxl.load_workbook(path,read_only=True)
    seen=0
    for row in book['Sheet'].values:
        expected=tuple((-1 if seen==0 and column==0 else seen*10+column) for column in range(10))
        assert row==expected
        seen+=1
    book.close();assert seen==rows

report={'rust':'1.88.0','openpyxl':'3.1.5','platform':platform.platform(),'cpu':next((line.split(':',1)[1].strip() for line in Path('/proc/cpuinfo').read_text().splitlines() if line.startswith('model name')),'unknown'),'measurement':'Native wait4 wall/CPU/RSS; one warmup and three Rust runs; one Python general-mode edit/save run without warmup; output/temp files polled at 25ms; checksum/interoperability excluded','scope':'M4 checkpoint, not complete existing-file structural editing; unchanged raw-copy, existing-cell overlays, cache invalidation, original assets and sparse core operations','cases':[]}
for rows in [10000,100000,1000000]:
    source=ROOT/f'benchmarks/data/numbers-{rows}.xlsx'
    if not source.exists():subprocess.run([sys.executable,str(ROOT/'benchmarks/write_baseline.py'),str(source),str(rows)],check=True)
    for mode in ['unchanged','edit']:
        runs=[]
        with tempfile.TemporaryDirectory(prefix='crabxl-edit-bench-') as temp:
            directory=Path(temp);output=directory/'out.xlsx'
            command=[ROOT/'target/release/examples/edit_demo',source,output,mode]
            measure(command,directory)
            for _ in range(3):
                result=measure(command,directory)
                copied,rewritten,xml_bytes,patch_bytes=map(int,result.pop('stdout').split())
                assert rewritten==(2 if mode=='edit' else 0)
                result.update(copied_parts=copied,rewritten_parts=rewritten,rewritten_xml_bytes=xml_bytes,charged_patch_bytes=patch_bytes,output_bytes=output.stat().st_size,logical_output_temp_bytes=output.stat().st_size)
                runs.append(result)
            checksum(output,rows,mode=='edit')
            if mode=='edit' and rows<=100000:public_readback(output,rows)
            if mode=='unchanged':
                with zipfile.ZipFile(source) as original,zipfile.ZipFile(output) as saved:
                    assert original.namelist()==saved.namelist()
                    assert all(original.getinfo(name).CRC==saved.getinfo(name).CRC and original.getinfo(name).compress_size==saved.getinfo(name).compress_size for name in original.namelist())
        case={'workload':mode,'rows':rows,'columns':10,'rust_runs':runs,'checksum_verified':True,'public_python_values_verified':mode=='edit' and rows<=100000}
        if mode=='edit' and rows<=100000:
            with tempfile.TemporaryDirectory(prefix='crabxl-python-edit-bench-') as temp:
                directory=Path(temp);output=directory/'python.xlsx'
                result=measure([sys.executable,ROOT/'benchmarks/edit_baseline.py',source,output],directory)
                assert result.pop('stdout')=='Edited A1';result['output_bytes']=output.stat().st_size
                checksum(output,rows,True);case['python_run']=result
        report['cases'].append(case)
        print('Preservation/edit',mode,rows,'checksum, output and cleanup passed',flush=True)
for rows in [10000,100000]:
    command=[ROOT/'target/release/examples/model_demo',rows];measure(command)
    runs=[]
    for _ in range(3):
        result=measure(command);cells,total,charged=map(int,result.pop('stdout').split());assert (cells,total)==(rows*10,(rows*10)*(rows*10-1)//2)
        result['charged_model_bytes']=charged;result['runtime_temporary_bytes']=0;runs.append(result)
    report['cases'].append({'workload':'sparse-model-build-insert-delete','rows':rows,'columns':10,'rust_runs':runs,'checksum_verified':True})
    print('Sparse model',rows,'structural checksum passed',flush=True)
for rows in [100000,1000000]:
    source=ROOT/f'benchmarks/data/numbers-{rows}.xlsx';commands={'M3':[args.before_reader,source],'M4':[ROOT/'target/release/examples/sum',source]}
    expected=measure(commands['M3'])['stdout'];assert measure(commands['M4'])['stdout']==expected
    runs=[]
    for index in range(3):
        for version in (['M3','M4'] if index%2==0 else ['M4','M3']):
            result=measure(commands[version]);assert result.pop('stdout')==expected;result['version']=version;runs.append(result)
    report['cases'].append({'workload':'reader-regression','rows':rows,'columns':10,'runs':runs,'checksum':expected,'before_revision':'4aaaee738f13','runtime_temporary_bytes':0})
    print('Reader regression',rows,'checksum passed',flush=True)
# Public generated workbook: styles, dates, comments, merges, dimensions, rules,
# image relationships and unknown original binary/XML parts survive two saves.
with tempfile.TemporaryDirectory(prefix='crabxl-interop-') as temp:
    directory=Path(temp);image=directory/'image.png';PillowImage.new('RGB',(8,8),'red').save(image)
    book=openpyxl.Workbook();sheet=book.active;sheet.title='Sheet';sheet['A1']=1;sheet['A1'].font=Font(bold=True,color='123456');sheet['B1']='=A1+1';sheet['C1']=datetime.datetime(2024,2,29);sheet['D1'].comment=Comment('Retain comment','Author');sheet['E1']='link';sheet['E1'].hyperlink='https://example.com';sheet.merge_cells('B3:C3');sheet.column_dimensions['A'].width=23;sheet.print_area='A1:E4';sheet.add_image(Image(image),'G1')
    rule=DataValidation(type='whole',operator='between',formula1='0',formula2='100');sheet.add_data_validation(rule);rule.add('A2:A4');book.properties.creator='Preservation fixture'
    original=directory/'original.xlsx';book.save(original);book.close()
    with zipfile.ZipFile(original) as archive:contents={name:archive.read(name) for name in archive.namelist()}
    contents['xl/worksheets/sheet1.xml']=contents['xl/worksheets/sheet1.xml'].replace(b'<f>A1+1</f><v></v>',b'<f>A1+1</f><v>2</v>')
    contents['custom/opaque.xml']=b'<unknown xmlns="urn:unknown"><![CDATA[untouched]]></unknown>'
    source=directory/'source.xlsx'
    with zipfile.ZipFile(source,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,data in contents.items():archive.writestr(name,data)
    output=directory/'saved.xlsx';command=[ROOT/'target/release/examples/edit_demo',source,output,'edit']
    for _ in range(2):
        subprocess.run(list(map(str,command)),check=True,capture_output=True)
        with zipfile.ZipFile(output) as archive:
            assert set(archive.namelist())==set(contents)
            for name,data in contents.items():
                if name not in ['xl/worksheets/sheet1.xml','xl/workbook.xml']:assert archive.read(name)==data,name
            assert archive.testzip() is None
        saved=openpyxl.load_workbook(output)
        active=saved['Sheet'];assert active['A1'].value==-1 and active['A1'].font.bold and active['A1'].font.color.rgb=='00123456'
        assert active['B1'].value=='=A1+1';assert active['C1'].value==datetime.datetime(2024,2,29);assert active['D1'].comment.text=='Retain comment';assert active['E1'].hyperlink.target=='https://example.com'
        assert str(active.merged_cells)=='B3:C3';assert active.column_dimensions['A'].width==23;assert active.data_validations.count==1;assert saved.properties.creator=='Preservation fixture';assert saved.calculation.fullCalcOnLoad and saved.calculation.forceFullCalc and saved.calculation.calcMode=='auto';saved.close()
        data=openpyxl.load_workbook(output,data_only=True);assert data['Sheet']['B1'].value is None;data.close()
    report['interoperability']={'public_generated_styles_dates_comments_hyperlinks_merges_dimensions_validation_images_properties':True,'unknown_part_bytes_and_unaffected_part_bytes_equal':True,'formula_cache_invalidated_and_recalc_requested':True,'repeated_save_verified':True}
    print('Public complex fixture, assets, properties and repeated saves passed',flush=True)
(ROOT/'benchmarks/results/m4-editor.json').write_text(json.dumps(report,indent=2)+'\n')
