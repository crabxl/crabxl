"""Measure sparse existing-package insertion without dense row expansion."""
import json
from pathlib import Path
import subprocess
import tempfile
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
report = {"baseline": "openpyxl 3.1.5 public readback", "scope": "Insert K1 and A1000001 into ten-column numeric files; no dense gap allocation", "measurement": "Native wait4 wall/CPU/RSS; one warmup and three runs; no worksheet XML temporary files; adjacent output ZIP space equals output bytes", "cases": []}
assert openpyxl.__version__ == "3.1.5"
for rows in (10000, 100000):
    source = ROOT / f"benchmarks/data/numbers-{rows}.xlsx"
    runs = []
    with tempfile.TemporaryDirectory(prefix="openrsxl-insert-") as name:
        output = Path(name) / "out.xlsx"
        command = [str(ROOT / "target/release/examples/edit_demo"), str(source), str(output), "insert"]
        subprocess.run(command, check=True, capture_output=True)
        for _ in range(3):
            run = subprocess.run([str(ROOT / "benchmarks/measure"), *command], check=True, capture_output=True, text=True)
            measured = json.loads(run.stderr.split("MEASURE ")[-1])
            copied, rewritten, rewritten_bytes, patch_bytes = map(int, run.stdout.split())
            measured.update(copied_parts=copied, rewritten_parts=rewritten, rewritten_bytes=rewritten_bytes, patch_bytes=patch_bytes, output_bytes=output.stat().st_size, logical_output_temp_bytes=output.stat().st_size, worksheet_temp_bytes=0)
            runs.append(measured)
        cells = rows * 10
        checksum = subprocess.run([str(ROOT / "target/release/examples/sum"), str(output)], check=True, capture_output=True, text=True).stdout.strip()
        assert checksum == f"{cells + 2} {cells * (cells - 1) // 2 + 16}"
        workbook = openpyxl.load_workbook(output)
        sheet = workbook["Sheet"]
        assert sheet["K1"].value == 7 and sheet["A1000001"].value == 9
        assert sheet.max_row == 1000001 and sheet.max_column == 11
        assert sheet["A1"].value == 0 and sheet.cell(rows, 10).value == cells - 1
        workbook.close()
        assert list(Path(name).iterdir()) == [output]
    report["cases"].append({"rows": rows, "runs": runs, "checksum": checksum, "public_readback": True, "cleanup": True})
    print(rows, "insertion, checksum, dimensions and cleanup passed", flush=True)
(ROOT / "benchmarks/results/m4-insertion.json").write_text(json.dumps(report, indent=2) + "\n")
