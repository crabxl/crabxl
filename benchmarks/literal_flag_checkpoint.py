"""Equivalent streaming raw data-table property verification."""
import json
import os
from pathlib import Path
import platform
import statistics
import sys
import tempfile
import openpyxl
from openpyxl.worksheet.formula import DataTableFormula
from iso_checkpoint import measure
LITERALS = ['', 'false', 'invalid', '0', ' true ']
FIELDS = ['ca', 'dt2D', 'dtr', 'del1', 'del2']
ROOT = Path(__file__).resolve().parents[1]
HERE = ROOT / 'benchmarks'

def generate(path, rows):
    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet('Sheet')
    for _ in range(rows):
        sheet.append([DataTableFormula('A1:B2', **dict.fromkeys(FIELDS, value)) for value in LITERALS])
    book.save(path)
    book.close()

def reference(path):
    book = openpyxl.load_workbook(path, read_only=True)
    count = 0
    for row in book.active:
        for column, cell in enumerate(row):
            value = LITERALS[column] or False
            assert all(getattr(cell.value, name) == value for name in FIELDS)
            count += 1
    book.close()
    print(count)

def main():
    target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
    report = {'reference': openpyxl.__version__, 'platform': platform.platform(), 'measurement': 'One warmup plus five rotating serial cold-process wall/CPU/RSS samples including process baseline; build/generation/cache readback excluded', 'semantics': 'Both readers stream and verify all five raw flag properties. Public defaults are False for omitted empty flags; canonical absent values remain absent and adapter supplies public constructor defaults. No prior core accepts every opaque flag, so no equivalent old baseline is fabricated. No temp storage used in this read workload.', 'cases': []}
    for rows in (1000, 10000):
        path = HERE / 'data' / f'literal-flags-{rows}.xlsx'
        generate(path, rows)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        assert all(cell.value is None for row in cached.active for cell in row)
        cached.close()
        commands = {'crabxl': [target / 'release/examples/literal_flag_read', path], 'openpyxl': [sys.executable, __file__, '--reference', path]}
        samples = {engine: [] for engine in commands}
        with tempfile.TemporaryDirectory() as name:
            temporary = Path(name)
            for command in commands.values():
                output, _ = measure(command, temporary)
                assert output == str(rows * 5)
            for iteration in range(5):
                for engine in (('crabxl', 'openpyxl') if iteration % 2 == 0 else ('openpyxl', 'crabxl')):
                    output, sample = measure(commands[engine], temporary)
                    assert output == str(rows * 5)
                    assert sample['sampled_temp_peak_bytes'] == 0
                    samples[engine].append(sample)
        report['cases'].append({'cells': rows * 5, 'input_bytes': path.stat().st_size, 'samples': samples, 'medians': {engine: {key: statistics.median(value[key] for value in values) for key in ('seconds', 'cpu_seconds', 'peak_rss_kib', 'sampled_temp_peak_bytes')} for engine, values in samples.items()}})
        (HERE / 'results/m2-literal-flags.json').write_text(json.dumps(report, indent=2)+'\n')
        print(rows, report['cases'][-1]['medians'], flush=True)

if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--reference':
        reference(sys.argv[2])
    else:
        main()
