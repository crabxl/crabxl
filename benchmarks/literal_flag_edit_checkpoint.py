"""Replace a known raw-flag formula using lazy preservation and public owned editing."""
import json
import os
from pathlib import Path
import platform
import statistics
import sys
import tempfile
import openpyxl
from openpyxl.worksheet.formula import DataTableFormula
from iso_checkpoint import measure
ROOT = Path(__file__).resolve().parents[1]
HERE = ROOT / 'benchmarks'
FIELDS = ['ca', 'dt2D', 'dtr', 'del1', 'del2']

def reference(source, output):
    book = openpyxl.load_workbook(source)
    book.active['A1'] = -1
    book.save(output)
    book.close()
    print('saved')

def verify(path, rows):
    book = openpyxl.load_workbook(path, read_only=True)
    for row in book.active:
        for cell in row:
            if cell.coordinate == 'A1':
                assert cell.value == -1
            else:
                assert cell.value.ref == 'A1:B2'
                assert all(getattr(cell.value, field) == 'opaque' for field in FIELDS)
    assert book.active.max_row in (None, rows)
    book.close()

def main():
    target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
    report = {'reference': openpyxl.__version__, 'platform': platform.platform(), 'measurement': 'One warmup plus five rotating serial cold-process wall/CPU/RSS samples including process baseline; generation and public verification excluded', 'semantics': 'Both replace A1 from a known data-table formula containing five opaque flags to -1 and save, retaining all other flag properties. Native edits lazily and preserves original parts; openpyxl explicitly owns a full editable model. This demonstrates feature/output overlap with different models, not identical mode or arbitrary graph editing. Output ZIP and native atomic ZIP staging live outside monitored TMPDIR; worksheet temporary samples are reported separately.', 'cases': []}
    with tempfile.TemporaryDirectory() as name:
        root = Path(name)
        temporary = root / 'temporary'
        temporary.mkdir()
        for rows in (1000, 10000):
            source = root / 'source.xlsx'
            book = openpyxl.Workbook(write_only=True)
            sheet = book.create_sheet('Sheet')
            for _ in range(rows):
                sheet.append([DataTableFormula('A1:B2', **dict.fromkeys(FIELDS, 'opaque')) for _ in range(5)])
            book.save(source)
            book.close()
            paths = {engine: root / f'{engine}.xlsx' for engine in ('crabxl', 'openpyxl')}
            commands = {'crabxl': [target / 'release/examples/edit_demo', source, paths['crabxl']], 'openpyxl': [sys.executable, __file__, '--reference', source, paths['openpyxl']]}
            samples = {engine: [] for engine in commands}
            for engine, command in commands.items():
                measure(command, temporary)
                verify(paths[engine], rows)
            for iteration in range(5):
                for engine in (('crabxl', 'openpyxl') if iteration % 2 == 0 else ('openpyxl', 'crabxl')):
                    _, sample = measure(commands[engine], temporary)
                    verify(paths[engine], rows)
                    sample['output_bytes'] = paths[engine].stat().st_size
                    samples[engine].append(sample)
            report['cases'].append({'cells': rows*5, 'samples': samples, 'medians': {engine: {key: statistics.median(value[key] for value in values) for key in ('seconds', 'cpu_seconds', 'peak_rss_kib', 'sampled_temp_peak_bytes')} for engine, values in samples.items()}})
            (HERE / 'results/m4-literal-flag-edit.json').write_text(json.dumps(report, indent=2)+'\n')
            print(rows, report['cases'][-1]['medians'], flush=True)
if __name__ == '__main__':
    if len(sys.argv) == 4 and sys.argv[1] == '--reference':
        reference(sys.argv[2], sys.argv[3])
    else:
        main()
