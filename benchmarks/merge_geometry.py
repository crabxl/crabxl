import importlib
import json
import resource
import sys
import time
import zipfile
from pathlib import Path

engine = importlib.import_module(sys.argv[1])
count = int(sys.argv[2])
book = engine.Workbook()
sheet = book.active
started = time.perf_counter()
for index in range(count):
    row = index * 3 + 1
    sheet.cell(row, 1, index)
    sheet.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=2)
prepared = time.perf_counter() - started
path = Path(f"/tmp/crabxl-w12-{sys.argv[1]}-{count}.xlsx")
started = time.perf_counter()
book.save(path)
saved = time.perf_counter() - started
peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
with zipfile.ZipFile(path) as archive:
    xml = archive.read("xl/worksheets/sheet1.xml")
    assert xml.count(b"<mergeCell ") == count
    assert b"A1:B2" in xml
import openpyxl

reopened = openpyxl.load_workbook(path)
assert len(reopened.active.merged_cells.ranges) == count
assert reopened.active.cell((count - 1) * 3 + 1, 1).value == count - 1
assert reopened.active.max_row == (count - 1) * 3 + 2
reopened.close()
book.close()
print(
    json.dumps(
        dict(
            engine=sys.argv[1],
            merges=count,
            prepare_seconds=prepared,
            save_seconds=saved,
            peak_rss_kib=peak,
            output_bytes=path.stat().st_size,
        )
    )
)
