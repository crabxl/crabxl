"""Compare compatible blank serialization with equivalent public readback."""
import argparse
import json
import os
from pathlib import Path
import platform
import statistics
import sys
import tempfile
import zipfile
import openpyxl
from shared_strings_checkpoint import ROOT, HERE, run
from iso_checkpoint import measure


def write_reference(path, count):
    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet("Sheet")
    for index in range(count):
        sheet.append([index + .25, float("inf"), -float("inf"), float("nan"), f'=A{index+1}+1'])
    book.save(path)
    book.close()
    print(count * 5)


def verify(path, count):
    for cached in (False, True):
        book = openpyxl.load_workbook(path, read_only=True, data_only=cached)
        seen = 0
        for index, row in enumerate(book.active.values):
            expected = (index + .25, None, None, None, None if cached else f'=A{index+1}+1')
            assert row == expected, (index, row)
            seen += 1
        assert seen == count
        book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows", type=int, nargs="+", default=[10000, 100000])
    parser.add_argument("--runs", type=int, default=5)
    parser.add_argument("--output", type=Path, default=HERE / "results/m2-nonfinite-write.json")
    args = parser.parse_args()
    assert openpyxl.__version__ == "3.1.5"
    if args.runs < 1 or any(not 0 < size <= 1048576 for size in args.rows):
        parser.error("Require positive samples and physical row bounds")
    run(["cargo", "build", "--release", "--locked", "-p", "crabxl", "--example", "nonfinite_write"])
    target = Path(os.environ.get("CARGO_TARGET_DIR", ROOT / "target"))
    report = {"reference": openpyxl.__version__, "platform": platform.platform(), "measurement": f"One warmup plus {args.runs} rotating serial wall/CPU/RSS samples; public normal/cache-only readback excluded from timing", "semantics": "Both sequential writers emit five columns: finite float, positive/negative infinity, NaN, formula with blank cache. Native receives an infinity formula cache while the public reference receives no cache; both serialize a blank value and reopen as None. Every reopened cell is checked in both modes. Native exact logical spool bytes include closed worksheet footer; 25ms temp sampling is a lower bound excluding final ZIP; no residual owned files. No calamine read or rust_xlsxwriter comparison is claimed for this compatibility checkpoint.", "cases": []}
    (HERE / "data").mkdir(exist_ok=True)
    for size in args.rows:
        paths = {name: HERE / "data" / f"nonfinite-{name}-{size}.xlsx" for name in ("crabxl", "openpyxl")}
        commands = {"crabxl": [target / "release/examples/nonfinite_write", paths["crabxl"], size], "openpyxl": [sys.executable, __file__, "--reference", paths["openpyxl"], size]}
        samples = {name: [] for name in commands}
        with tempfile.TemporaryDirectory() as directory:
            for name, command in commands.items():
                output, _ = measure(command, Path(directory))
                assert output == str(size * 5)
                verify(paths[name], size)
            names = list(commands)
            for iteration in range(args.runs):
                for name in names[iteration % 2:] + names[:iteration % 2]:
                    output, sample = measure(commands[name], Path(directory))
                    assert output == str(size * 5)
                    verify(paths[name], size)
                    with zipfile.ZipFile(paths[name]) as archive:
                        sample["worksheet_xml_bytes"] = archive.getinfo("xl/worksheets/sheet1.xml").file_size
                    sample["output_bytes"] = paths[name].stat().st_size
                    if name == "crabxl":
                        assert sample["logical_temp_peak_bytes"] == sample["worksheet_xml_bytes"]
                    samples[name].append(sample)
                    print(size, name, sample, flush=True)
        report["cases"].append({"rows": size, "cells": size*5, "samples": samples, "medians": {name: {key: statistics.median(s[key] for s in values) for key in ("seconds", "cpu_seconds", "peak_rss_kib", "sampled_temp_peak_bytes")} for name, values in samples.items()}})
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(report, indent=2) + "\n")


if __name__ == "__main__":
    if len(sys.argv) == 4 and sys.argv[1] == "--reference":
        write_reference(sys.argv[2], int(sys.argv[3]))
    else:
        main()
