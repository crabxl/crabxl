"""Equivalent requested-prefix formulas; native additionally validates the XML tail."""
import argparse
import io
import json
import os
from pathlib import Path
import platform
import statistics
import sys
import tempfile
import zipfile
import openpyxl
from shared_strings_checkpoint import ROOT, HERE, run
from iso_checkpoint import measure
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

def generate(path, rows):
    source = io.BytesIO()
    openpyxl.Workbook().save(source)
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            if name != "xl/worksheets/sheet1.xml":
                archive.writestr(name, value)
        with archive.open("xl/worksheets/sheet1.xml", "w") as sheet:
            sheet.write(f'<worksheet xmlns="{MAIN}"><sheetData>'.encode())
            for index in range(1, rows + 1):
                sheet.write(f'<row r="{index}"><c r="A{index}"><f t="shared" si="{index}" ref="A{index}">A{index}+1</f><v>{index}</v></c></row>'.encode())
            sheet.write(b'</sheetData></worksheet>')

def reference(path, limit):
    book = openpyxl.load_workbook(path, read_only=True)
    count = 0
    for index, row in enumerate(book.active.iter_rows(max_row=limit), 1):
        assert row[0].value == f"=A{index}+1"
        count += 1
    book.close()
    print(count)

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", type=Path, required=True)
    parser.add_argument("--baseline-core", required=True)
    parser.add_argument("--runs", type=int, default=5)
    args = parser.parse_args()
    assert openpyxl.__version__ == "3.1.5"
    target = Path(os.environ.get("CARGO_TARGET_DIR", ROOT / "target"))
    report = {"reference": openpyxl.__version__, "baseline_core": args.baseline_core, "platform": platform.platform(), "measurement": f"One warmup and {args.runs} rotating serial samples including process baseline, no builds/checks during timing", "semantics": "All readers verify the same 1000 requested formula expressions. Native also checks every integer cache and consumes the entire XML/CRC tail; public cache-only validation is a separate untimed pass, and its projected iterator stops before validating the tail. Previous engine uses its direct ReadOptions stream because a composable joint-policy API did not exist; current uses a 2MiB joint managed budget. The supplied files have valid tails. Neither native retains worksheet XML/cells or uses temporary storage. Old retention of later unrelated masters is measured explicitly. Reference speed remains a cold process workload result; parser-only equal validation is not claimed.", "cases": []}
    (HERE / "data").mkdir(exist_ok=True)
    for total in (10000, 100000):
        limit = 1000
        path = HERE / "data" / f"projected-shared-{total}.xlsx"
        generate(path, total)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        assert [row[0].value for row in cached.active.iter_rows(max_row=limit)] == list(range(1, limit+1))
        cached.close()
        commands = {"crabxl": [target / "release/examples/policy_projected_formulas", path, limit], "crabxl-before": [args.baseline.resolve(), path, limit], "openpyxl": [sys.executable, __file__, "--reference", path, limit]}
        template_counts = {}
        for name in ("crabxl", "crabxl-before"):
            result = run(commands[name])
            template_counts[name] = int(next(line.split()[1] for line in result.stderr.splitlines() if line.startswith("PROJECTED_STATS ")))
        assert template_counts == {"crabxl": limit, "crabxl-before": total}
        samples = {name: [] for name in commands}
        with tempfile.TemporaryDirectory() as directory:
            for command in commands.values():
                output, _ = measure(command, Path(directory))
                assert output == str(limit)
            names = list(commands)
            for iteration in range(args.runs):
                for name in names[iteration % 3:] + names[:iteration % 3]:
                    output, sample = measure(commands[name], Path(directory))
                    assert output == str(limit)
                    assert sample["sampled_temp_peak_bytes"] == 0
                    samples[name].append(sample)
                    print(total, name, sample, flush=True)
        report["cases"].append({"source_rows": total, "requested_rows": limit, "templates": template_counts, "input_bytes": path.stat().st_size, "samples": samples, "medians": {name: {key: statistics.median(sample[key] for sample in values) for key in ("seconds", "cpu_seconds", "peak_rss_kib", "sampled_temp_peak_bytes")} for name, values in samples.items()}})
        (HERE / "results/m2-policy-options.json").write_text(json.dumps(report, indent=2) + "\n")

if __name__ == "__main__":
    if len(sys.argv) == 4 and sys.argv[1] == "--reference":
        reference(sys.argv[2], int(sys.argv[3]))
    else:
        main()
