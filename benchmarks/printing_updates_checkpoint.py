"""Repeated margin updates with large unrelated page-break lists."""
import argparse
import hashlib
import json
from pathlib import Path
import statistics
import sys
import tempfile
import openpyxl
from openpyxl.worksheet.pagebreak import Break
from iso_checkpoint import measure
from printing_checkpoint import signature
from worksheet_feature_checkpoint import verify_values


def reference(source, target, updates):
    book = openpyxl.load_workbook(source)
    for index in range(updates):
        book.active.page_margins.left = index / 100.0
    book.save(target)
    book.close()
    print(updates)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--baseline', type=Path, required=True)
    parser.add_argument('--target', type=Path, required=True, help='Current release/examples directory')
    parser.add_argument('--output', type=Path, default=Path(__file__).parent / 'results/m5-printing-updates.json')
    args = parser.parse_args()
    assert openpyxl.__version__ == '3.1.5'
    updates = 20
    report = {'baseline_core': '8136b1d837e8f5ef6e67f32d2a1effbe7c9d2341', 'baseline_binary_sha256': hashlib.sha256(args.baseline.read_bytes()).hexdigest(), 'reference': openpyxl.__version__, 'measurement': 'One warmup and five rotating serial Linux wait4 samples including startup; fixture generation, builds/tests, public property and every-cell verification excluded. Temporary files sampled every 25ms with cleanup checked.', 'semantics': 'Twenty left-margin changes with unrelated source page-break lists, followed by one save. Prior/current snapshot workers use identical public Rust calls; component worker uses checked component moves. Reference uses ordinary load/property/save calls. This intentionally stresses unrelated metadata ownership and repeated source validation, not every printing workload. Pure metadata edits retain numeric content.', 'cases': []}
    with tempfile.TemporaryDirectory() as name:
        root = Path(name)
        temporary = root / 'temporary'
        temporary.mkdir()
        for rows, breaks in ((5000, 1000), (50000, 10000)):
            source = root / 'source.xlsx'
            book = openpyxl.Workbook(write_only=True)
            sheet = book.create_sheet('Sheet')
            sheet.row_breaks.brk = [Break(id=index + 1) for index in range(breaks)]
            for row in range(rows):
                sheet.append([row])
            book.save(source)
            book.close()
            expected_path = root / 'expected.xlsx'
            reference(source, expected_path, updates)
            expected = signature(expected_path)
            paths = {engine: root / f'{engine}.xlsx' for engine in ('prior-snapshot', 'current-snapshot', 'current-component', 'openpyxl')}
            commands = {
                'prior-snapshot': [args.baseline.resolve(), source, paths['prior-snapshot'], updates],
                'current-snapshot': [args.target.resolve() / 'printing_snapshot_updates', source, paths['current-snapshot'], updates],
                'current-component': [args.target.resolve() / 'printing_component_updates', source, paths['current-component'], updates],
                'openpyxl': [sys.executable, __file__, '--reference', source, paths['openpyxl'], updates],
            }
            samples = {engine: [] for engine in commands}
            for engine, command in commands.items():
                output, _ = measure(command, temporary)
                assert output.splitlines()[-1] == str(updates)
                assert signature(paths[engine]) == expected
                verify_values(paths[engine], rows)
            engines = list(commands)
            for iteration in range(5):
                for engine in engines[iteration % 4:] + engines[:iteration % 4]:
                    output, sample = measure(commands[engine], temporary)
                    assert output.splitlines()[-1] == str(updates)
                    for line in output.splitlines():
                        if line.startswith('PRINT_BYTES='):
                            sample['managed_overlay_bytes'] = int(line.split('=')[1])
                    assert signature(paths[engine]) == expected
                    verify_values(paths[engine], rows)
                    sample['output_bytes'] = paths[engine].stat().st_size
                    samples[engine].append(sample)
            report['cases'].append({'rows': rows, 'breaks': breaks, 'updates': updates, 'source_bytes': source.stat().st_size, 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'samples': samples, 'medians': {engine: {key: statistics.median(item[key] for item in values) for key in ('seconds', 'cpu_seconds', 'peak_rss_kib', 'sampled_temp_peak_bytes')} for engine, values in samples.items()}})
            args.output.write_text(json.dumps(report, indent=2) + '\n')


if __name__ == '__main__':
    if len(sys.argv) > 1 and sys.argv[1] == '--reference':
        reference(sys.argv[2], sys.argv[3], int(sys.argv[4]))
    else:
        main()
