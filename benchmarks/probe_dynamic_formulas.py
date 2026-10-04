"""Public cm/vm projection observations plus native read/edit interoperability."""
import argparse
import io
import json
from pathlib import Path
import subprocess
import tempfile
import zipfile
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--reader", type=Path, required=True)
parser.add_argument("--editor", type=Path, required=True)
args = parser.parse_args()
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
output = io.BytesIO()
openpyxl.Workbook().save(output)
with zipfile.ZipFile(output) as archive:
    base = {name: archive.read(name) for name in archive.namelist()}
checked = []
with tempfile.TemporaryDirectory() as directory:
    source = Path(directory) / "dynamic.xlsx"
    target = Path(directory) / "edited.xlsx"
    for token in ["1", "0", "-1", "4294967295", "not-an-index"]:
        parts = dict(base)
        parts["xl/worksheets/sheet1.xml"] = f'<worksheet xmlns="{MAIN}"><sheetData><row r="1"><c r="A1" cm="{token}"><f t="array" ref="A1">_xlfn.SEQUENCE(1)</f><v>1</v></c><c r="B1" vm="{token}"><v>1</v></c></row></sheetData></worksheet>'.encode()
        parts["xl/metadata.xml"] = f'<metadata xmlns="{MAIN}"><extLst><ext uri="opaque-source"/></extLst></metadata>'.encode()
        with zipfile.ZipFile(source, "w") as archive:
            for name, value in parts.items():
                archive.writestr(name, value)
        book = openpyxl.load_workbook(source)
        formula = book.active["A1"].value
        assert isinstance(formula, ArrayFormula)
        assert (formula.text, formula.ref) == ("=_xlfn.SEQUENCE(1)", "A1")
        assert book.active["B1"].value == 1
        cached = openpyxl.load_workbook(source, data_only=True)
        assert [cached.active.cell(1, i).value for i in (1, 2)] == [1, 1]
        assert subprocess.check_output([str(args.reader.resolve()), str(source)], text=True).strip() == "2"
        subprocess.run([str(args.editor.resolve()), str(source), str(target), "edit"], check=True, capture_output=True)
        native_edited = openpyxl.load_workbook(target)
        assert [native_edited.active.cell(1, i).value for i in (1, 2)] == [-1, 1]
        with zipfile.ZipFile(target) as archive:
            assert archive.read("xl/metadata.xml") == parts["xl/metadata.xml"]
            xml = archive.read("xl/worksheets/sheet1.xml").decode()
            assert 'cm=' not in xml and f'vm="{token}"' in xml
        checked.append({"source_metadata": token, "array_expression": formula.text, "array_range": formula.ref, "cached_values": [1, 1], "native_read": "checked", "native_target_replacement": "checked", "unaffected_metadata_part": "preserved"})
root = Path(__file__).resolve().parents[1]
report = {"reference": openpyxl.__version__, "inspection": "Public APIs and generated OOXML only; opaque indices are not interpreted", "checked": checked}
(root / "docs/research/dynamic-formulas-public.json").write_text(json.dumps(report, indent=2) + "\n")
(root / "benchmarks/results/m2-dynamic-formulas-interop.json").write_text(json.dumps(report, indent=2) + "\n")
