"""Public constructor/save/reload and canonical reader style-domain assertions."""
import json
import subprocess
from pathlib import Path
import os
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.styles import Font, GradientFill, Color

ROOT = Path(__file__).resolve().parents[1]
path = ROOT / "benchmarks/data/style-domains.xlsx"
path.parent.mkdir(exist_ok=True)
book = openpyxl.Workbook()
for row, size in enumerate((-1.0, 410.0, 1_000_000.0), 1):
    cell = book.active.cell(row, 1, row)
    cell.font = Font(sz=size)
book.active['A1'].fill = GradientFill(type='path', left=-1, right=2)
book.active['A1'].number_format = ''
color = Color(rgb='invalid', theme=1, indexed=3, auto=False)
assert color.type == 'indexed' and color.value == 3
book.active['A2'].font = Font(sz=410, color=color)
book.save(path)
book.close()
loaded = openpyxl.load_workbook(path)
assert [loaded.active.cell(row, 1).font.sz for row in range(1, 4)] == [-1.0, 410.0, 1_000_000.0]
assert loaded.active['A1'].fill.left == -1 and loaded.active['A1'].fill.right == 2
assert loaded.active['A1'].number_format == ''
assert loaded.active['A2'].font.color.type == 'indexed'
assert loaded.active['A2'].font.color.indexed == 3
loaded.close()
# The new-package style exporter stages table-style schemas. Normalize only this
# generated fixture's zero-record declarations; no source preservation is claimed.
with zipfile.ZipFile(path) as archive:
    parts = {name: archive.read(name) for name in archive.namelist()}
root = ET.fromstring(parts['xl/styles.xml'])
for node in list(root):
    if node.tag.rsplit('}', 1)[-1] in ('tableStyles', 'dxfs'):
        assert len(node) == 0 and int(node.get('count', 0)) == 0
        root.remove(node)
parts['xl/styles.xml'] = ET.tostring(root)
with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
    for name, value in parts.items():
        archive.writestr(name, value)
target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
native = path.with_name('style-domains-native.xlsx')
output = subprocess.check_output([target / 'release/examples/style_domain_probe', path, native], text=True).strip()
assert output == 'finite-style-domains-ok'
loaded = openpyxl.load_workbook(native)
assert [loaded.active.cell(row, 1).font.sz for row in range(1, 4)] == [-1.0, 410.0, 1_000_000.0]
assert loaded.active['A1'].fill.left == -1 and loaded.active['A1'].fill.right == 2
assert loaded.active['A1'].number_format == ''
assert loaded.active['A2'].font.color.type == 'indexed' and loaded.active['A2'].font.color.indexed == 3
loaded.close()
report = {'reference': openpyxl.__version__, 'font_sizes': [-1.0, 410.0, 1_000_000.0], 'gradient_edges': [-1.0, 2.0], 'number_format': '', 'color_selection': 'indexed before theme, auto and rgb', 'public_save_reload': 'pass', 'canonical_reader': output}
(ROOT / 'benchmarks/results/m2-finite-style-domains-interop.json').write_text(json.dumps(report, indent=2) + '\n')
print(json.dumps(report))
