"""Compare newly assigned data-table properties through public save/reload calls."""
import argparse
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile
import openpyxl
from openpyxl.worksheet.formula import DataTableFormula

def inspect(book):
    return {cell: vars(book.active[cell].value) for cell in ('D1', 'I1')}

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('native', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    reference = openpyxl.Workbook()
    reference.active['D1'] = DataTableFormula(ref='D1:E2', dt2D=True, dtr=False, r1='$A$1', r2='B1', del1=False, del2=True)
    reference.active['I1'] = DataTableFormula(ref='I1:J2', dt2D=False, dtr=False, r1='')
    before = inspect(reference)
    output = BytesIO()
    reference.save(output)
    reference.close()
    output.seek(0)
    reloaded = openpyxl.load_workbook(output)
    expected = inspect(reloaded)
    native = openpyxl.load_workbook(args.native)
    actual = inspect(native)
    assert expected == actual, (expected, actual)
    with ZipFile(args.native) as archive:
        xml = archive.read('xl/worksheets/sheet1.xml').decode()
    args.output.write_text(json.dumps({'reference': openpyxl.__version__, 'inspection': 'Public constructors and generated save/reload only; no implementation inspection', 'before': before, 'reference_reloaded': expected, 'native_reloaded': actual, 'native_xml': xml}, indent=2) + '\n')
    reloaded.close()
    native.close()

if __name__ == '__main__':
    main()
