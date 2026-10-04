"""Public visible normal/shared text when source hints are not used."""
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import zipfile
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
attributes = [
    't="future" ref="invalid"',
    'ca="invalid" si="invalid" r1="opaque" t="normal"',
    'unknown="value" t="normal"',
    'ca="invalid" unknown="value" si="3" ref="opaque" t="shared"',
]
book = openpyxl.Workbook()
stream = BytesIO()
book.save(stream)
book.close()
with zipfile.ZipFile(stream) as archive:
    parts = {name: archive.read(name) for name in archive.namelist()}
results = []
for index, hints in enumerate(attributes):
    path = ROOT / f'benchmarks/data/formula-header-{index}.xlsx'
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            if name == 'xl/worksheets/sheet1.xml':
                payload = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1"><f '+hints+'>A1+1</f><v>2</v></c></row></sheetData></worksheet>').encode()
            archive.writestr(name, payload)
    loaded = openpyxl.load_workbook(path)
    assert loaded.active['A1'].value == '=A1+1'
    loaded.close()
    cached = openpyxl.load_workbook(path, data_only=True)
    assert cached.active['A1'].value == 2
    cached.close()
    assert subprocess.check_output([target / 'release/examples/formula_header_read', path], text=True).strip() == '1'
    results.append({'source_attributes': hints, 'formula': '=A1+1', 'cache': 2, 'native_verified': True})
report = {'reference': openpyxl.__version__, 'cases': results, 'scope': 'Visible normal/shared formulas only; structured array/table raw flag properties remain separate work. Typed compatible reads discard unused source hints; original-package preservation remains separate.'}
(ROOT / 'benchmarks/results/m2-formula-header-interop.json').write_text(json.dumps(report, indent=2)+'\n')
print(json.dumps(report))
