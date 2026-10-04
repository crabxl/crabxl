"""Public raw data-table flag properties and canonical XML interoperability."""
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.worksheet.formula import DataTableFormula
ROOT = Path(__file__).resolve().parents[1]
MAIN = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
LITERALS = ['', 'false', 'invalid', '0', ' true ']
FIELDS = ['ca', 'dt2D', 'dtr', 'del1', 'del2']
target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
path = ROOT / 'benchmarks/data/literal-flags-native.xlsx'
assert subprocess.check_output([target / 'release/examples/literal_flag_fixture', path], text=True).strip() == '5'
assert subprocess.check_output([target / 'release/examples/literal_flag_read', path], text=True).strip() == '5'
book = openpyxl.Workbook()
for column, literal in enumerate(LITERALS, 1):
    book.active.cell(1, column, DataTableFormula('A1:B2', **dict.fromkeys(FIELDS, literal)))
    assert all(getattr(book.active.cell(1, column).value, name) == literal for name in FIELDS)
stream = BytesIO()
book.save(stream)
book.close()
def records(source):
    with zipfile.ZipFile(source) as archive:
        root = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        return [dict(node.attrib) for node in root.iter(MAIN+'f')]
assert records(path) == records(stream)
checked = []
for source in (path, stream):
    stream.seek(0)
    loaded = openpyxl.load_workbook(source)
    for column, literal in enumerate(LITERALS, 1):
        value = loaded.active.cell(1, column).value
        assert all(getattr(value, name) == (literal if literal else False) for name in FIELDS)
    checked.append([vars(loaded.active.cell(1, column).value) for column in range(1, 6)])
    loaded.close()
report = {'reference': openpyxl.__version__, 'literals': LITERALS, 'fields': FIELDS, 'xml_records': records(path), 'public_readback': checked}
(ROOT / 'benchmarks/results/m2-literal-flags-interop.json').write_text(json.dumps(report, indent=2)+'\n')
print(json.dumps(report))
