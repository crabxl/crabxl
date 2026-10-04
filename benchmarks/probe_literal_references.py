"""Public literal formula properties and canonical XML interoperability."""
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

ROOT = Path(__file__).resolve().parents[1]
MAIN = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
path = ROOT / 'benchmarks/data/literal-reference-native.xlsx'
target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
assert subprocess.check_output([target / 'release/examples/literal_reference_fixture', path], text=True).strip() == '9'
references = ['', 'A1:B2', '$A$1:$B$2', 'Sheet1!A1:B2', 'not a range']
inputs = ['', 'C1', 'Sheet1!C1', 'input']
book = openpyxl.Workbook()
for column, value in enumerate(references, 1):
    book.active.cell(1, column, ArrayFormula(value, '=1'))
for column, value in enumerate(inputs, 6):
    book.active.cell(1, column, DataTableFormula('opaque', r1=value, r2=''))
assert [book.active.cell(1, c).value.ref for c in range(1, 6)] == references
assert [book.active.cell(1, c).value.r1 for c in range(6, 10)] == inputs
stream = BytesIO()
book.save(stream)
book.close()
def records(source):
    with zipfile.ZipFile(source) as archive:
        root = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        return [{'attributes': dict(n.attrib), 'body': n.text} for n in root.iter(MAIN+'f')]
assert records(path) == records(stream)
properties = []
for source in (path, stream):
    stream.seek(0)
    loaded = openpyxl.load_workbook(source)
    assert [loaded.active.cell(1, c).value.ref for c in range(1, 6)] == [None] + references[1:]
    assert [loaded.active.cell(1, c).value.r1 for c in range(6, 10)] == [None, 'C1', 'Sheet1!C1', 'input']
    assert [loaded.active.cell(1, c).value.r2 for c in range(6, 10)] == [None]*4
    properties.append([vars(loaded.active.cell(1, c).value) for c in range(1, 10)])
    loaded.close()
    stream.seek(0)
    cached = openpyxl.load_workbook(source, data_only=True)
    assert [cached.active.cell(1, c).value for c in range(1, 10)] == [None]*9
    cached.close()
report = {'reference': openpyxl.__version__, 'xml_records': records(path), 'public_readback': properties, 'cache_projection': [None]*9}
(ROOT / 'benchmarks/results/m2-literal-references-interop.json').write_text(json.dumps(report, indent=2)+'\n')
print(json.dumps(report))
