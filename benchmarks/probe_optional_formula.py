"""Public optional array expression properties and canonical XML interoperability."""
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

ROOT = Path(__file__).resolve().parents[1]
MAIN = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
path = ROOT / 'benchmarks/data/optional-array-native.xlsx'
target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
assert subprocess.check_output([target / 'release/examples/optional_formula_fixture', path], text=True).strip() == '3'
reference = openpyxl.Workbook()
inputs = [None, '', '=1']
for column, text in enumerate(inputs, 1):
    reference.active.cell(1, column, ArrayFormula(None, text))
assert [reference.active.cell(1, column).value.text for column in range(1, 4)] == inputs
stream = BytesIO()
reference.save(stream)
reference.close()
def records(source):
    with zipfile.ZipFile(source) as archive:
        root = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        return [{'attributes': dict(node.attrib), 'body': node.text} for node in root.iter(MAIN+'f')]
assert records(path) == records(stream)
checked = []
for source in (path, stream):
    stream.seek(0)
    book = openpyxl.load_workbook(source)
    assert [vars(book.active.cell(1, col).value) for col in range(1, 4)] == [{'ref': None, 'text': '='}, {'ref': None, 'text': '='}, {'ref': None, 'text': '=1'}]
    checked.append([vars(book.active.cell(1, col).value) for col in range(1, 4)])
    book.close()
    stream.seek(0)
    cached = openpyxl.load_workbook(source, data_only=True)
    assert [cached.active.cell(1, col).value for col in range(1, 4)] == [None] * 3
    cached.close()
# Reference creation permits absent data-table ref, but ordinary reload fails.
# Record the public defect without reproducing a crash in canonical I/O.
table = openpyxl.Workbook()
table.active['A1'] = DataTableFormula(None)
output = BytesIO()
table.save(output)
table.close()
output.seek(0)
try:
    openpyxl.load_workbook(output)
except TypeError as error:
    table_defect = {'exception': type(error).__name__, 'message': str(error)}
else:
    raise AssertionError('Expected pinned reference missing-ref defect')
report = {'reference': openpyxl.__version__, 'literal_expression_presence': inputs, 'xml_records_equal': records(path), 'public_readback': checked, 'cache_projection': [None]*3, 'missing_data_table_ref_reference_defect': table_defect}
(ROOT / 'benchmarks/results/m2-optional-formula-interop.json').write_text(json.dumps(report, indent=2)+'\n')
print(json.dumps(report))
