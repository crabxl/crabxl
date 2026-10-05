"""Literal/missing shared identifiers compared through public formula loading."""
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import tempfile
import zipfile
from xml.sax.saxutils import quoteattr
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
CASES = [(value, value) for value in (None, '', '0', '01', '1', '+1', '-1', '4294967296', 'text', '1 ', 'group & value')] + [('01', '1'), ('1', '01'), (None, ''), ('', None)]

def main():
    target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
    workbook = openpyxl.Workbook()
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    with zipfile.ZipFile(stream) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    results = []
    with tempfile.TemporaryDirectory() as name:
        for master, follower in CASES:
            path = Path(name) / 'source.xlsx'
            def attribute(value):
                return '' if value is None else ' si=' + quoteattr(value)
            xml = '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1"><f t="shared"' + attribute(master) + '>A1+1</f><v>2</v></c></row><row r="2"><c r="A2"><f t="shared"' + attribute(follower) + '/><v>3</v></c></row></sheetData></worksheet>'
            with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
                for part, payload in parts.items():
                    archive.writestr(part, xml.encode() if part == 'xl/worksheets/sheet1.xml' else payload)
            loaded = openpyxl.load_workbook(path)
            expected = [loaded.active.cell(row, 1).value for row in (1, 2)]
            loaded.close()
            native = subprocess.check_output([target / 'release/examples/shared_index_read', path], text=True).splitlines()
            assert native == expected
            cached = openpyxl.load_workbook(path, data_only=True)
            assert [cached.active.cell(row, 1).value for row in (1, 2)] == [2, 3]
            cached.close()
            results.append({'master': master, 'follower': follower, 'visible_formulas': expected, 'caches': [2, 3], 'native_matches': True})
    report = {'reference': openpyxl.__version__, 'cases': results, 'scope': 'Compatible reading and identifier spelling; strict numeric schema validation and group editing remain separate'}
    (ROOT / 'benchmarks/results/m2-shared-index-interop.json').write_text(json.dumps(report, indent=2) + '\n')
    print(f'{len(results)} shared-index cases match public openpyxl')
if __name__ == '__main__':
    main()
