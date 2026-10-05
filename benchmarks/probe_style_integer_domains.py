"""Public exact style identities and the reference's large-integer save limitation."""
import argparse
from io import BytesIO
import json
from pathlib import Path
import subprocess
from zipfile import ZipFile
import openpyxl
from openpyxl.styles import Color, Font

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--native', type=Path, required=True)
parser.add_argument('--output', type=Path, required=True)
args = parser.parse_args()
assert openpyxl.__version__ == '3.1.5'
with __import__('tempfile').TemporaryDirectory() as directory:
    fixture = Path(directory) / 'integers.xlsx'
    sizes = subprocess.check_output([str(args.native.resolve()), str(fixture)], text=True).strip()
    book = openpyxl.load_workbook(fixture)
    checked = []
    reference_saves = []
    for row, (number, kind) in enumerate([(n, k) for n in [2**63, -2**63-1, 10**40] for k in ['theme', 'indexed']], 1):
        cell = book.active.cell(row, 1)
        assert cell.value == 1 and cell.font.charset == number
        for color in [cell.font.color, cell.fill.fgColor, cell.border.left.color]:
            assert color.type == kind and getattr(color, kind) == number
        checked.append({'cell': cell.coordinate, 'charset': str(number), 'kind': kind, 'identity': str(number)})
        reference = openpyxl.Workbook()
        reference.active['A1'] = 1
        reference.active['A1'].font = Font(charset=number, color=Color(**{kind: number}))
        assert reference.active['A1'].font.charset == number
        buffer = BytesIO()
        reference.save(buffer)
        buffer.seek(0)
        with ZipFile(buffer) as archive:
            style_xml = archive.read('xl/styles.xml').decode()
        buffer.seek(0)
        try:
            openpyxl.load_workbook(buffer)
        except TypeError as error:
            reference_saves.append({'integer': str(number), 'kind': kind, 'reload_error': str(error), 'scientific_notation': 'e+' in style_xml})
        else:
            raise AssertionError('Expected the pinned reference save/reload limitation')
    book.close()
    report = {'reference': openpyxl.__version__, 'construction_and_native_output_readback': checked, 'reference_self_save': reference_saves, 'native_sizes': dict(zip(['StyleInteger', 'Color', 'Font', 'catalog_bytes'], map(int, sizes.split()))), 'scope': 'Exact native output and imported style identities; no theme/palette resolution or arbitrary-precision Python style proxy is claimed.'}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2) + '\n')
