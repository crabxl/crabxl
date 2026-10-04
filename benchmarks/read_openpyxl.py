"""Numeric read-only baseline; process startup and workbook discovery included."""
import sys
from openpyxl import load_workbook

workbook = load_workbook(sys.argv[1], read_only=True, data_only=True)
try:
    count = 0
    checksum = 0
    for row in workbook["Sheet"].iter_rows(values_only=True):
        for value in row:
            checksum += value
            count += 1
    print(count, checksum)
finally:
    workbook.close()
