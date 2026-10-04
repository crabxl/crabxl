"""Equivalent expanded-expression workloads; native additionally validates every cache."""
import argparse
import json
import os
from pathlib import Path
import platform
import statistics
import sys
import tempfile
import zipfile
import openpyxl
from shared_strings_checkpoint import ROOT, HERE, run
from iso_checkpoint import measure


def generate(path, rows):
    book = openpyxl.Workbook()
    book.save(path)
    book.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            if name != "xl/worksheets/sheet1.xml":
                archive.writestr(name, value)
        with archive.open("xl/worksheets/sheet1.xml", "w") as sheet:
            sheet.write(b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>')
            for index in range(1, rows + 1):
                shared = f'<f t="shared" si="4294967295" ref="A1:A{rows}">A1+$Z$1</f>' if index == 1 else '<f t="shared" si="4294967295"/>'
                sheet.write(f'<row r="{index}"><c r="A{index}">{shared}<v>{index}</v></c><c r="B{index}"><f>B{index}+1</f><v>{index}</v></c></row>'.encode())
            sheet.write(b'</sheetData></worksheet>')


def reference(path):
    book = openpyxl.load_workbook(path, read_only=True)
    count = 0
    for index, row in enumerate(book.active, 1):
        assert len(row) == 2
        assert row[0].value == f'=A{index}+$Z$1'
        assert row[1].value == f'=B{index}+1'
        count += 2
    book.close()
    print(count)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows", type=int, nargs="+", default=[10000, 100000])
    parser.add_argument("--baseline", type=Path, help="Optional exact prior native binary for rotating regression samples")
    parser.add_argument("--runs", type=int, default=5)
    parser.add_argument("--output", type=Path, default=HERE / "results/m2-shared-formulas.json")
    args = parser.parse_args()
    assert openpyxl.__version__ == "3.1.5"
    if args.runs < 1 or any(not 0 < size <= 1048576 for size in args.rows):
        parser.error("Require positive samples and physical worksheet row bounds")
    run(["cc", "-O2", "-Wall", "-Wextra", "-Werror", HERE / "measure.c", "-o", HERE / "measure"])
    run(["cargo", "build", "--release", "--locked", "-p", "crabxl", "--example", "shared_formula_read"])
    target = Path(os.environ.get("CARGO_TARGET_DIR", ROOT / "target"))
    report = {"reference": openpyxl.__version__, "platform": platform.platform(), "measurement": f"One warmup and {args.runs} rotating serial wall/CPU/RSS samples including process baseline", "semantics": "Both readers stream and verify every expanded formula. Native also verifies every cached integer in the same pass; public openpyxl cache verification uses a separate untimed data_only pass. Shared template accounting is a conservative managed-storage estimate, not exact RSS. No temporary storage is used for this read workload. calamine comparison is deferred until equivalent streaming formula/cache APIs are validated.", "cases": []}
    (HERE / "data").mkdir(exist_ok=True)
    for size in args.rows:
        path = HERE / "data" / f"shared-formulas-{size}.xlsx"
        generate(path, size)
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        checked = 0
        for index, row in enumerate(book.active, 1):
            assert [cell.value for cell in row] == [index, index]
            checked += 2
        book.close()
        assert checked == size * 2
        commands = {"crabxl": [target / "release/examples/shared_formula_read", path], "openpyxl": [sys.executable, __file__, "--reference", path]}
        if args.baseline:
            commands["crabxl-before"] = [args.baseline.resolve(), path]
        samples = {name: [] for name in commands}
        with tempfile.TemporaryDirectory() as directory:
            for command in commands.values():
                output, _ = measure(command, Path(directory))
                assert output == str(checked)
            names = list(commands)
            for iteration in range(args.runs):
                for name in names[iteration % len(names):] + names[:iteration % len(names)]:
                    output, sample = measure(commands[name], Path(directory))
                    assert output == str(checked)
                    assert sample["sampled_temp_peak_bytes"] == 0
                    samples[name].append(sample)
                    print(size, name, sample, flush=True)
        native = run(commands["crabxl"])
        stats_line = next(line for line in native.stderr.splitlines() if line.startswith("FORMULA_STATS "))
        templates, byte_count, expanded, unresolved = map(int, stats_line.split()[1:])
        assert (templates, expanded, unresolved) == (1, size - 1, 0)
        report["cases"].append({"rows": size, "cells": checked, "input_bytes": path.stat().st_size, "shared_templates": templates, "accounted_template_bytes": byte_count, "samples": samples, "medians": {name: {key: statistics.median(s[key] for s in values) for key in ("seconds", "cpu_seconds", "peak_rss_kib", "sampled_temp_peak_bytes")} for name, values in samples.items()}})
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(report, indent=2) + "\n")


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--reference":
        reference(sys.argv[2])
    else:
        main()
