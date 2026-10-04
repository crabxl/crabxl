"""Equivalent source styled-value model edit and save, with public readback."""
import json
import os
from pathlib import Path
import platform
import statistics
import sys
import tempfile
import zipfile
import openpyxl
from shared_strings_checkpoint import ROOT, HERE, run
from iso_checkpoint import measure

def reference(input_path, output_path):
    book = openpyxl.load_workbook(input_path, data_only=True)
    book.active['A1'].number_format = 'mm-dd-yy'
    count = sum(len(row) for row in book.active.iter_rows())
    book.save(output_path)
    book.close()
    print(count)

def properties(path):
    book = openpyxl.load_workbook(path, data_only=True)
    cells = []
    for row in book.active.iter_rows():
        for cell in row:
            value = cell.value
            if hasattr(value, 'isoformat'): value = value.isoformat()
            elif hasattr(value, 'total_seconds'): value = value.total_seconds()
            cells.append({'value': value, 'number_format': cell.number_format, 'font_name': cell.font.name, 'font_size': cell.font.sz, 'font_family': cell.font.family, 'font_scheme': cell.font.scheme, 'font_bold': cell.font.b, 'fill_pattern': cell.fill.patternType, 'alignment': {key: getattr(cell.alignment, key) for key in ('horizontal', 'vertical', 'textRotation', 'wrapText', 'shrinkToFit', 'indent')}, 'protection': {'locked': cell.protection.locked, 'hidden': cell.protection.hidden}})
    book.close()
    return cells

def main():
    target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
    report = {'reference': openpyxl.__version__, 'platform': platform.platform(), 'measurement': 'One warmup and five rotating serial cold-process samples, no builds during timing', 'semantics': 'Both engines fully load ten stored typed/cache values and declared style catalogs, change A1 number format, and save a new package. All saved values and listed public style properties are compared outside timing. Source fixtures contain no advanced/opaque part graphs; original-package preservation is not claimed. Native source styles retain IDs and automatically registered date formats derive from source-zero component IDs. Native additionally retains unused source declarations and four automatic date presets; public save prunes unused codes. Output size/table counts expose that extra native work. Previously no equivalent source-style export API existed.', 'cases': []}
    for extra in (1000, 50000):
        input_path = HERE / 'data' / f'style-import-{extra}.xlsx'
        assert input_path.exists(), 'Run style_import_checkpoint.py to generate input fixtures'
        outputs = {name: HERE / 'data' / f'style-export-{name}-{extra}.xlsx' for name in ('crabxl', 'openpyxl')}
        commands = {'crabxl': [target / 'release/examples/style_catalog_export', input_path, outputs['crabxl']], 'openpyxl': [sys.executable, __file__, '--reference', input_path, outputs['openpyxl']]}
        samples = {name: [] for name in commands}
        with tempfile.TemporaryDirectory() as directory:
            for command in commands.values():
                output, _ = measure(command, Path(directory)); assert output == '10'
            expected = properties(outputs['openpyxl'])
            assert properties(outputs['crabxl']) == expected
            names = list(commands)
            for iteration in range(5):
                for name in names[iteration % 2:] + names[:iteration % 2]:
                    output, sample = measure(commands[name], Path(directory)); assert output == '10'
                    assert properties(outputs[name]) == expected
                    sample['output_bytes'] = outputs[name].stat().st_size
                    samples[name].append(sample)
        temporary = int(run(commands['crabxl']).stderr.split('STYLE_EXPORT_TEMP ')[1].strip())
        tables = {}
        import xml.etree.ElementTree as ET
        for name, path in outputs.items():
            with zipfile.ZipFile(path) as archive:
                raw = archive.read('xl/styles.xml')
                root = ET.fromstring(raw)
                tables[name] = {'number_formats': len(root.find('{*}numFmts')), 'cell_formats': len(root.find('{*}cellXfs')), 'style_xml_bytes': len(raw)}
        report['cases'].append({'declared_formats': extra+1, 'cells': 10, 'output_tables': tables, 'native_logical_temp_bytes': temporary, 'verified_public_properties': expected, 'samples': samples, 'medians': {name: {key: statistics.median(sample[key] for sample in values) for key in ('seconds', 'cpu_seconds', 'peak_rss_kib', 'sampled_temp_peak_bytes', 'output_bytes')} for name, values in samples.items()}})
        (HERE / 'results/m2-style-export.json').write_text(json.dumps(report, indent=2)+'\n')
        print(json.dumps(report['cases'][-1]['medians']), flush=True)

if __name__ == '__main__':
    if len(sys.argv) == 4 and sys.argv[1] == '--reference': reference(sys.argv[2], sys.argv[3])
    else: main()
