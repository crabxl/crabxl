"""Typed differential/table definition transfer through generated public fixtures."""
import json
import os
from pathlib import Path
import platform
import statistics
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.styles.differential import DifferentialStyle, DifferentialStyleList
from openpyxl.styles.numbers import NumberFormat
from openpyxl.styles.table import TableStyleList, TableStyle, TableStyleElement
from openpyxl.xml.functions import tostring
from shared_strings_checkpoint import ROOT, HERE
from iso_checkpoint import measure
MAIN = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REGIONS = sorted(TableStyleElement.type.values)

def namespaced(tree):
    value = ET.fromstring(tostring(tree))
    for node in value.iter():
        if not node.tag.startswith('{'):
            node.tag = MAIN + node.tag
    return value

def generate(path, count):
    book = openpyxl.Workbook()
    book.active['A1'] = 1
    book.save(path)
    book.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(parts['xl/styles.xml'])
    for node in list(root):
        if node.tag in (MAIN+'dxfs', MAIN+'tableStyles'):
            root.remove(node)
    differentials = ET.Element(MAIN+'dxfs', count='4000000000')
    for index in range(count):
        value = DifferentialStyle(font=Font(name=f'Diff{index}', b=True, color='FFAABBCC'), numFmt=NumberFormat(numFmtId=500, formatCode='0.000'), fill=PatternFill(patternType='solid', fgColor='FF123456'), alignment=Alignment(horizontal='right', indent=2.5, wrap_text=True), border=Border(left=Side(style='thin', color='FF102030')), protection=Protection(locked=False, hidden=True))
        differentials.append(namespaced(value.to_tree()))
    tables = TableStyleList(defaultTableStyle='Custom & Named', defaultPivotStyle='PivotStyleLight16', tableStyle=[TableStyle(name='Custom & Named', pivot=False, table=True, count=123, tableStyleElement=[TableStyleElement(type=region, size=0 if index == 0 else 2, dxfId=index % count) for index, region in enumerate(REGIONS)])])
    root.append(differentials)
    root.append(namespaced(tables.to_tree()))
    parts['xl/styles.xml'] = ET.tostring(root)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)

def verify(path, count):
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read('xl/styles.xml'))
    differentials = DifferentialStyleList.from_tree(root.find(MAIN+'dxfs'))
    assert len(differentials.dxf) == count
    for index, value in enumerate(differentials.dxf):
        assert value.font.name == f'Diff{index}' and value.font.b is True
        assert value.font.color.rgb == 'FFAABBCC'
        assert value.numFmt.numFmtId == 500 and value.numFmt.formatCode == '0.000'
        assert value.fill.patternType == 'solid' and value.fill.fgColor.rgb == 'FF123456'
        assert value.alignment.horizontal == 'right' and value.alignment.indent == 2.5 and value.alignment.wrap_text is True
        assert value.border.left.style == 'thin' and value.border.left.color.rgb == 'FF102030'
        assert value.protection.locked is False and value.protection.hidden is True
    tables = TableStyleList.from_tree(root.find(MAIN+'tableStyles'))
    assert tables.defaultTableStyle == 'Custom & Named' and tables.defaultPivotStyle == 'PivotStyleLight16'
    assert len(tables.tableStyle) == 1
    style = tables.tableStyle[0]
    assert style.name == 'Custom & Named' and style.count == 123 and style.pivot is False and style.table is True
    assert [(value.type, value.size, value.dxfId) for value in style.tableStyleElement] == [(region, 0 if index == 0 else 2, index % count) for index, region in enumerate(REGIONS)]
    book = openpyxl.load_workbook(path)
    assert book.active['A1'].value == 1
    book.close()

def reference(source, output, count):
    book = openpyxl.load_workbook(source)
    book.save(output)
    book.close()
    print(count)

def main():
    assert openpyxl.__version__ == '3.1.5'
    target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
    report = {'reference': openpyxl.__version__, 'platform': platform.platform(), 'measurement': 'One warmup and five rotating serial process wall/CPU/RSS samples; builds/fixture generation/public readback excluded', 'semantics': 'Read and new-package export of unused differential/table catalog definitions plus one scalar cell. All differential properties, 28 table/pivot regions, references, defaults and optional declared count are checked outside timing through public metadata classes and workbook scalar access. Reference loads/saves its workbook; native streams a row and adopts source catalogs, additionally verifying each differential font identity. Source-style IDs are stable. This does not create or preserve worksheet conditional-formatting/table instances or advanced package graphs. No earlier equivalent native differential export API exists; no prior comparison is fabricated.', 'cases': []}
    for count in (1000, 10000):
        source = HERE / 'data' / f'style-extras-{count}.xlsx'
        generate(source, count)
        verify(source, count)
        paths = {name: HERE / 'data' / f'style-extras-{name}-{count}.xlsx' for name in ('crabxl', 'openpyxl')}
        commands = {'crabxl': [target / 'release/examples/style_extras_export', source, paths['crabxl'], count], 'openpyxl': [sys.executable, __file__, '--reference', source, paths['openpyxl'], count]}
        samples = {name: [] for name in commands}
        with tempfile.TemporaryDirectory() as directory:
            for name, command in commands.items():
                output, _ = measure(command, Path(directory))
                assert output == str(count)
                verify(paths[name], count)
            for iteration in range(5):
                for name in (('crabxl', 'openpyxl') if iteration % 2 == 0 else ('openpyxl', 'crabxl')):
                    output, sample = measure(commands[name], Path(directory))
                    assert output == str(count)
                    verify(paths[name], count)
                    sample['output_bytes'] = paths[name].stat().st_size
                    samples[name].append(sample)
        stats = subprocess.run(list(map(str, commands['crabxl'])), text=True, capture_output=True, check=True)
        managed = int(next(line.split()[1] for line in stats.stderr.splitlines() if line.startswith('STYLE_BYTES ')))
        metadata_bytes = {}
        for name, path in paths.items():
            with zipfile.ZipFile(path) as archive:
                metadata_bytes[name] = archive.getinfo('xl/styles.xml').file_size
        report['cases'].append({'differentials': count, 'native_managed_style_bytes': managed, 'style_xml_bytes': metadata_bytes, 'table_regions': len(REGIONS), 'source_bytes': source.stat().st_size, 'samples': samples, 'medians': {name: {key: statistics.median(sample[key] for sample in values) for key in ('seconds', 'cpu_seconds', 'peak_rss_kib', 'sampled_temp_peak_bytes')} for name, values in samples.items()}})
        (HERE / 'results/m2-style-extras.json').write_text(json.dumps(report, indent=2)+'\n')
        print(count, report['cases'][-1]['medians'], flush=True)

if __name__ == '__main__':
    if len(sys.argv) > 1 and sys.argv[1] == '--reference':
        reference(Path(sys.argv[2]), Path(sys.argv[3]), int(sys.argv[4]))
    else:
        main()
