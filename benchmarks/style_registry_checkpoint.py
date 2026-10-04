"""Shared component catalog creation with verified combination identities and values."""
import argparse
import json
import os
from pathlib import Path
import platform
import statistics
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from shared_strings_checkpoint import ROOT, HERE, run
from iso_checkpoint import measure
MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def reference(path, count):
    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet("Sheet")
    fonts = [Font(name="Calibri", size=11, family=2, scheme="minor", color=f'{0xFF000000 | index:08X}') for index in range(16)]
    fills = [PatternFill(patternType="solid", fgColor=f'{0xFF100000 | index:08X}') for index in range(16)]
    alignments = [Alignment(horizontal="general", vertical="bottom", textRotation=index, wrap_text=True, shrink_to_fit=True) for index in range(32)]
    for index in range(count):
        cell = WriteOnlyCell(sheet, value=index+.25)
        cell.font, cell.fill, cell.alignment = fonts[index % 16], fills[(index // 16) % 16], alignments[(index // 256) % 32]
        cell.number_format = "0.000"
        sheet.append([cell])
    book.save(path)
    book.close()
    print(count)


def verify(path, count):
    book = openpyxl.load_workbook(path, read_only=True)
    seen = 0
    for index, row in enumerate(book.active):
        assert len(row) == 1
        cell = row[0]
        assert cell.value == index+.25
        assert cell.number_format == "0.000"
        assert cell.font.name == "Calibri" and cell.font.sz == 11
        assert cell.font.color.rgb == f'{0xFF000000 | (index % 16):08X}'
        assert cell.fill.patternType == "solid"
        assert cell.fill.fgColor.rgb == f'{0xFF100000 | ((index // 16) % 16):08X}'
        assert cell.alignment.textRotation == (index // 256) % 32
        assert cell.alignment.horizontal == "general" and cell.alignment.vertical == "bottom"
        assert cell.alignment.wrap_text is True and cell.alignment.shrink_to_fit is True
        assert cell.protection.locked is True and cell.protection.hidden is False
        seen += 1
    assert seen == count
    book.close()
    with zipfile.ZipFile(path) as archive:
        styles = ET.fromstring(archive.read("xl/styles.xml"))
        counts = {name: len(styles.find(MAIN+name)) for name in ("fonts", "fills", "borders", "numFmts", "cellXfs")}
        counts["style_xml_bytes"] = archive.getinfo("xl/styles.xml").file_size
        counts["worksheet_xml_bytes"] = archive.getinfo("xl/worksheets/sheet1.xml").file_size
        return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--styles", type=int, nargs="+", default=[1000, 8000])
    parser.add_argument("--runs", type=int, default=5)
    parser.add_argument("--baseline", type=Path, required=True)
    parser.add_argument("--baseline-core", default="3dfda8a5b7d3fb376d781cc74cad84867a34e4a5")
    parser.add_argument("--output", type=Path, default=HERE / "results/m2-style-registry.json")
    args = parser.parse_args()
    assert openpyxl.__version__ == "3.1.5"
    if args.runs < 1 or any(not 0 < count <= 8192 for count in args.styles):
        parser.error("Require positive samples and 1..8192 combinations")
    target = Path(os.environ.get("CARGO_TARGET_DIR", ROOT / "target"))
    report = {"reference": openpyxl.__version__, "platform": platform.platform(), "baseline_core": args.baseline_core, "measurement": f"One warmup plus {args.runs} rotating serial Linux wall/CPU/RSS samples including process baseline. Builds and public/native readback excluded from creation timing.", "semantics": "All writers create one floating cell per distinct font/fill/rotation combination using sequential spools. Both native versions additionally re-register every style to assert stable IDs; public reference shares component objects and interns formats during row writing. Timed workflows produce equivalent public values/styles, not identical API call counts. Every cell is checked through both public and canonical native streaming readers outside timing. Source font/fill/format counts and metadata bytes are recorded separately from RAM; native exact completed spool and lower-bound 25ms temp sampling exclude final ZIP. No calamine creation overlap or rust_xlsxwriter API comparison is claimed here.", "cases": []}
    (HERE / "data").mkdir(exist_ok=True)
    for count in args.styles:
        paths = {name: HERE / "data" / f"style-registry-{name}-{count}.xlsx" for name in ("crabxl", "crabxl-before", "openpyxl")}
        commands = {"crabxl": [target / "release/examples/style_registry_fixture", paths["crabxl"], count], "crabxl-before": [args.baseline.resolve(), paths["crabxl-before"], count], "openpyxl": [sys.executable, __file__, "--reference", paths["openpyxl"], count]}
        samples = {name: [] for name in commands}
        tables = {}
        with tempfile.TemporaryDirectory() as directory:
            for name, command in commands.items():
                output, _ = measure(command, Path(directory))
                assert output == str(count)
                tables[name] = verify(paths[name], count)
                assert run([target / "release/examples/style_registry_read", paths[name]]).stdout.strip() == str(count)
            names = list(commands)
            for iteration in range(args.runs):
                for name in names[iteration % 3:] + names[:iteration % 3]:
                    output, sample = measure(commands[name], Path(directory))
                    assert output == str(count)
                    assert verify(paths[name], count) == tables[name]
                    assert run([target / "release/examples/style_registry_read", paths[name]]).stdout.strip() == str(count)
                    sample["output_bytes"] = paths[name].stat().st_size
                    if "logical_temp_peak_bytes" in sample:
                        assert sample["logical_temp_peak_bytes"] == tables[name]["worksheet_xml_bytes"]
                    samples[name].append(sample)
                    print(count, name, sample, flush=True)
        fields = list(map(int, run([target / "release/examples/style_registry_fixture", "unused", count, "stats"]).stdout.split()))
        assert fields[:5] == [17, 18, 1, 5, count+5]
        report["cases"].append({"styles": count, "cells": count, "tables": tables, "native_accounted_style_bytes": fields[5], "samples": samples, "medians": {name: {key: statistics.median(s[key] for s in values) for key in ("seconds", "cpu_seconds", "peak_rss_kib", "sampled_temp_peak_bytes")} for name, values in samples.items()}})
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(report, indent=2)+"\n")


if __name__ == "__main__":
    if len(sys.argv) == 4 and sys.argv[1] == "--reference":
        reference(sys.argv[2], int(sys.argv[3]))
    else:
        main()
