"""Streaming temporal assignment preserves components and existing date formats."""
import datetime
import json
import os
from pathlib import Path
import platform
import statistics
import sys
import tempfile
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from iso_checkpoint import measure
ROOT = Path(__file__).resolve().parents[1]
HERE = ROOT / 'benchmarks'
FORMATS = ['General', '0.00', 'yyyy-mm-dd', 'hh:mm:ss', '[h]:mm:ss']
DATES = [datetime.date(2024, 1, 2), datetime.datetime(2024, 1, 2, 3, 4, 5, 678900), datetime.time(3, 4, 5, 678900), datetime.timedelta(days=2, seconds=3, microseconds=678900)]

def generate(path, rows, streaming=True):
    book = openpyxl.Workbook(write_only=streaming)
    sheet = book.create_sheet('Sheet') if streaming else book.active
    font = Font(name='Public temporal font', sz=14, b=True, color='FF112233')
    fill = PatternFill(patternType='solid', fgColor='FF445566')
    border = Border(left=Side(style='thin', color='FF778899'))
    alignment = Alignment(indent=2.5, wrap_text=True)
    for row in range(1, rows+1):
        cells = []
        for column, code in enumerate(FORMATS, 1):
            cell = WriteOnlyCell(sheet) if streaming else sheet.cell(row, column)
            cell.font, cell.fill, cell.border, cell.alignment = font, fill, border, alignment
            cell.number_format = code
            cell.value = DATES[(row-1) % len(DATES)]
            cells.append(cell)
        if streaming:
            sheet.append(cells)
    book.save(path)
    book.close()

def record(cell):
    return {'value_type': type(cell.value).__name__, 'value': str(cell.value), 'format': cell.number_format,
            'font': [cell.font.name, cell.font.sz, cell.font.b, cell.font.color.type, cell.font.color.rgb],
            'fill': [cell.fill.patternType, cell.fill.fgColor.type, cell.fill.fgColor.rgb],
            'border': [cell.border.left.style, cell.border.left.color.type, cell.border.left.color.rgb],
            'alignment': [cell.alignment.indent, cell.alignment.wrap_text]}

def records(path):
    book = openpyxl.load_workbook(path, read_only=True)
    values = [record(cell) for row in book.active for cell in row]
    book.close()
    return values

def verify(path, rows, expected):
    book = openpyxl.load_workbook(path, read_only=True)
    count = 0
    for row in book.active:
        for cell in row:
            assert record(cell) == expected[count % len(expected)]
            count += 1
    book.close()
    assert count == rows*5

def main():
    target = Path(os.environ.get('CARGO_TARGET_DIR', ROOT / 'target'))
    report = {'reference': openpyxl.__version__, 'platform': platform.platform(), 'measurement': 'One warmup plus five rotating serial cold-process wall/CPU/RSS samples; process/import baseline included; generation of expected properties/build/readback excluded', 'semantics': 'Both writers stream five formats across four temporal kinds. Each output value/type, format, font/fill/border and alignment is checked using public readback. Native source styles register once and derive shared formats; public WriteOnlyCell assignments set the format before the temporal value. No earlier native API supports all explicit non-date/different-kind assignments, so no equivalent old feature baseline is fabricated.', 'cases': []}
    with tempfile.TemporaryDirectory() as name:
        root = Path(name)
        temporary = root / 'temporary'
        temporary.mkdir()
        reference = root / 'reference.xlsx'
        generate(reference, 4, streaming=False)
        expected = records(reference)
        for rows in (1000, 10000):
            paths = {engine: root / f'{engine}.xlsx' for engine in ('crabxl', 'openpyxl')}
            commands = {'crabxl': [target / 'release/examples/temporal_style_fixture', paths['crabxl'], rows], 'openpyxl': [sys.executable, __file__, '--reference', paths['openpyxl'], rows]}
            samples = {engine: [] for engine in commands}
            for engine, command in commands.items():
                output, _ = measure(command, temporary)
                assert output == str(rows*5)
                verify(paths[engine], rows, expected)
            for iteration in range(5):
                for engine in (('crabxl', 'openpyxl') if iteration % 2 == 0 else ('openpyxl', 'crabxl')):
                    output, sample = measure(commands[engine], temporary)
                    assert output == str(rows*5)
                    verify(paths[engine], rows, expected)
                    sample['output_bytes'] = paths[engine].stat().st_size
                    samples[engine].append(sample)
            report['cases'].append({'cells': rows*5, 'samples': samples, 'medians': {engine: {key: statistics.median(sample[key] for sample in values) for key in ('seconds', 'cpu_seconds', 'peak_rss_kib', 'sampled_temp_peak_bytes')} for engine, values in samples.items()}})
            (HERE / 'results/m2-temporal-styles.json').write_text(json.dumps(report, indent=2)+'\n')
            print(rows, report['cases'][-1]['medians'], flush=True)
if __name__ == '__main__':
    if len(sys.argv) == 4 and sys.argv[1] == '--reference':
        generate(Path(sys.argv[2]), int(sys.argv[3]))
        print(int(sys.argv[3])*5)
    else:
        main()
