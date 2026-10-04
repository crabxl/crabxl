"""Verify exact opaque/default themes using public APIs and generated ZIP data."""
import argparse
import hashlib
import io
import json
from pathlib import Path
import subprocess
import tempfile
import zipfile
import openpyxl
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--native", type=Path, required=True)
args = parser.parse_args()
reference = io.BytesIO()
openpyxl.Workbook().save(reference)
with zipfile.ZipFile(reference) as archive:
    default = archive.read("xl/theme/theme1.xml")
custom = b'<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="custom"><a:extLst/></a:theme>'
checked = []
with tempfile.TemporaryDirectory() as directory:
    for mode, expected in [("default", default), ("empty", default), ("custom", custom), ("opaque", b"not XML"), ("omit", None)]:
        path = Path(directory) / f"{mode}.xlsx"
        stats = json.loads(subprocess.check_output([str(args.native.resolve()), str(path), mode], text=True))
        book = openpyxl.load_workbook(path)
        assert book.active["A1"].value == 42
        assert book.loaded_theme == expected
        assert stats["theme_bytes"] == (len(expected) if expected else 0)
        checked.append({"mode": mode, **stats, "sha256": hashlib.sha256(expected).hexdigest() if expected else None})
    observed = []
    for value in [b"not XML", b"<wrong/>", b"", b"<!DOCTYPE x><x/>"]:
        book = openpyxl.Workbook()
        book.loaded_theme = value
        output = io.BytesIO()
        book.save(output)
        loaded = openpyxl.load_workbook(io.BytesIO(output.getvalue())).loaded_theme
        assert loaded == (value or default)
        observed.append({"input": value.decode(), "accepted": True, "uses_default": not value})
root = Path(__file__).resolve().parents[1]
(root / "benchmarks/results/m2-theme-interop.json").write_text(json.dumps({"reference": openpyxl.__version__, "checked": checked, "public_opaque_observations": observed}, indent=2) + "\n")
