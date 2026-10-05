"""Shared serial harness for canonical worksheet metadata ports and public probes."""
import json
import os
from pathlib import Path
import statistics
import subprocess
import sys
import tempfile
import openpyxl
from iso_checkpoint import measure

ROOT = Path(__file__).resolve().parents[1]
HERE = ROOT / "benchmarks"


def python_operation(mode, source, output, rows, configure):
    if mode == "create":
        book = openpyxl.Workbook(write_only=True)
        sheet = book.create_sheet("Sheet")
        configure(sheet)
        for row in range(rows): sheet.append([row])
    elif mode == "edit":
        book = openpyxl.load_workbook(source)
        configure(book["Sheet"])
    else:
        raise ValueError(mode)
    book.save(output)
    book.close()


def verify_values(path, rows):
    book = openpyxl.load_workbook(path, read_only=True)
    count = 0
    for count, values in enumerate(book["Sheet"].iter_rows(values_only=True), 1):
        assert values == (count - 1,), (path, count, values)
    assert count == rows
    book.close()


def run_checkpoint(args, scenario):
    """Scenario supplies public configuration/signature and native metadata telemetry."""
    configure, signature = scenario["configure"], scenario["signature"]
    if args.python_operation:
        python_operation(args.python_operation, args.source, args.target, args.rows[0], configure)
        return
    assert openpyxl.__version__ == "3.1.5"
    assert args.runs > 0 and all(0 < rows <= 1048576 for rows in args.rows)
    target = Path(os.environ.get("CARGO_TARGET_DIR", ROOT / "target")) / "release/examples" / scenario["example"]
    report = {"reference": openpyxl.__version__, "measurement": "One warmup plus rotating serial Linux wait4 samples; builds and all-cell/public-metadata verification outside timing", "semantics": scenario["semantics"], "cases": []}
    data = HERE / "data"
    data.mkdir(exist_ok=True)
    expected_path = data / f"{scenario['name']}-public-expected.xlsx"
    python_operation("create", None, expected_path, 1, configure)
    expected = signature(expected_path)
    for rows in args.rows:
        source = data / f"{scenario['name']}-source-{rows}.xlsx"
        book = openpyxl.Workbook(write_only=True)
        sheet = book.create_sheet("Sheet")
        for row in range(rows): sheet.append([row])
        book.save(source)
        for mode in ("create", "edit"):
            paths = {name: data / f"{scenario['name']}-{mode}-{name}-{rows}.xlsx" for name in ("crabxl", "openpyxl")}
            commands = {"crabxl": [target, mode, source, paths["crabxl"], rows], "openpyxl": [sys.executable, scenario["script"], "--python-operation", mode, "--source", source, "--target", paths["openpyxl"], "--rows", rows]}
            samples = {name: [] for name in commands}
            with tempfile.TemporaryDirectory(prefix=f"{scenario['name']}-bench-") as directory:
                for command in commands.values(): measure(command, Path(directory))
                for iteration in range(args.runs):
                    names = list(commands)
                    for name in names[iteration % 2:] + names[:iteration % 2]:
                        output, sample = measure(commands[name], Path(directory))
                        for line in output.splitlines():
                            if line.startswith("TEMP_BYTES="): sample["logical_temp_peak_bytes"] = int(line.split("=")[1])
                            if line.startswith(scenario["telemetry"] + "="): sample[scenario["managed_key"]] = int(line.split("=")[1])
                        assert signature(paths[name]) == expected, paths[name]
                        verify_values(paths[name], rows)
                        samples[name].append(sample)
                copied = data / f"{scenario['name']}-copy-{rows}.xlsx"
                subprocess.run(list(map(str, [target, "copy", paths["openpyxl"], copied, rows])), check=True, env=dict(os.environ, TMPDIR=directory), capture_output=True)
                assert signature(copied) == expected
                verify_values(copied, rows)
                assert not list(Path(directory).iterdir())
            case = {"rows": rows, "operation": mode, "samples": samples, "medians": {name: {key: statistics.median(sample[key] for sample in values) for key in ("seconds", "cpu_seconds", "peak_rss_kib", "sampled_temp_peak_bytes")} for name, values in samples.items()}}
            report["cases"].append(case)
            print(rows, mode, case["medians"], flush=True)
    args.output.write_text(json.dumps(report, indent=2) + "\n")
