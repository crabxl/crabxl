"""Serial creation and original-package view editing with public feature verification."""
import argparse
import json
import os
from pathlib import Path
import statistics
import subprocess
import sys
import tempfile
import zipfile
import openpyxl
from openpyxl.worksheet.views import SheetView, SheetViewList, Pane
from iso_checkpoint import measure

ROOT = Path(__file__).resolve().parents[1]
HERE = ROOT / "benchmarks"


def settings(sheet):
    sheet.freeze_panes = "B3"
    view = sheet.sheet_view
    for name, value in dict(windowProtection=False, showFormulas=True, showGridLines=False,
                            showRowColHeaders=False, showZeros=False, rightToLeft=True,
                            tabSelected=True, showRuler=False, showOutlineSymbols=False,
                            defaultGridColor=False, showWhiteSpace=False, zoomToFit=True,
                            view="pageLayout", topLeftCell="B3", colorId=12, zoomScale=145,
                            zoomScaleNormal=80, zoomScaleSheetLayoutView=90,
                            zoomScalePageLayoutView=95).items():
        setattr(view, name, value)
    view.selection[2].activeCellId = 2
    view.selection[2].sqref = "A1 C3:D4"
    sheet.views.sheetView.append(SheetView(workbookViewId=2, pane=Pane(
        xSplit=-1.5, ySplit=2.5, topLeftCell="opaque", activePane="bottomRight", state="frozenSplit")))


def python_operation(mode, source, output, rows):
    if mode == "create":
        book = openpyxl.Workbook(write_only=True)
        sheet = book.create_sheet("Sheet")
        settings(sheet)
        for row in range(rows):
            sheet.append([row])
    elif mode == "edit":
        book = openpyxl.load_workbook(source)
        settings(book["Sheet"])
    else:
        raise ValueError(mode)
    book.save(output)
    book.close()


def signature(views):
    """Public descriptor values and source order, with load-time defaults applied."""
    records = []
    for view in views.sheetView:
        record = {key: value for key, value in vars(view).items() if key not in ("pane", "selection")}
        record["pane"] = None if view.pane is None else vars(view.pane)
        record["selection"] = [vars(value) for value in view.selection]
        records.append(record)
    return records


def verify(path, rows, expected):
    with zipfile.ZipFile(path) as archive:
        from openpyxl.xml.functions import fromstring
        from xml.etree import ElementTree as ET
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        element = root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheetViews")
        actual = SheetViewList.from_tree(fromstring(ET.tostring(element)))
        assert signature(actual) == expected, (path, signature(actual), expected)
    book = openpyxl.load_workbook(path, read_only=True)
    count = 0
    for count, values in enumerate(book["Sheet"].iter_rows(values_only=True), 1):
        assert values == (count - 1,), (path, count, values)
    assert count == rows
    book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rows", type=int, nargs="+", default=[5000, 50000])
    parser.add_argument("--runs", type=int, default=5)
    parser.add_argument("--output", type=Path, default=HERE / "results/m5-worksheet-views.json")
    parser.add_argument("--python-operation", choices=["create", "edit"])
    parser.add_argument("--source", type=Path)
    parser.add_argument("--target", type=Path)
    args = parser.parse_args()
    if args.python_operation:
        python_operation(args.python_operation, args.source, args.target, args.rows[0])
        return
    assert openpyxl.__version__ == "3.1.5"
    target = Path(os.environ.get("CARGO_TARGET_DIR", ROOT / "target")) / "release/examples/worksheet_views"
    report = {"reference": openpyxl.__version__, "measurement": "One warmup plus rotating serial Linux wait4 samples; builds and all-cell/public-view verification outside timing", "semantics": "Create: sequential spools in both engines. Edit: native bounded original-package overlay, reference ordinary loaded workbook. Both output identical public view properties and every numeric value. No claim of equivalent ownership cost. Read/export interoperability checked outside timing. 25ms temporary sampling is a lower bound and excludes final ZIP; native creation reports exact worksheet spool peak. No upstream comparison: multiple views and arbitrary public attributes are not yet mapped to an equivalent upstream public workload.", "cases": []}
    data = HERE / "data"
    data.mkdir(exist_ok=True)
    expected_path = data / "views-public-expected.xlsx"
    python_operation("create", None, expected_path, 1)
    expected = signature(openpyxl.load_workbook(expected_path).active.views)
    for rows in args.rows:
        source = data / f"views-source-{rows}.xlsx"
        book = openpyxl.Workbook(write_only=True)
        sheet = book.create_sheet("Sheet")
        for row in range(rows): sheet.append([row])
        book.save(source)
        for mode in ("create", "edit"):
            paths = {name: data / f"views-{mode}-{name}-{rows}.xlsx" for name in ("crabxl", "openpyxl")}
            commands = {"crabxl": [target, mode, source, paths["crabxl"], rows], "openpyxl": [sys.executable, Path(__file__).resolve(), "--python-operation", mode, "--source", source, "--target", paths["openpyxl"], "--rows", rows]}
            samples = {name: [] for name in commands}
            with tempfile.TemporaryDirectory(prefix="views-bench-") as directory:
                for command in commands.values(): measure(command, Path(directory))
                for iteration in range(args.runs):
                    names = list(commands)
                    for name in names[iteration % 2:] + names[:iteration % 2]:
                        output, sample = measure(commands[name], Path(directory))
                        for line in output.splitlines():
                            if line.startswith("TEMP_BYTES="): sample["logical_temp_peak_bytes"] = int(line.split("=")[1])
                            if line.startswith("VIEW_BYTES="): sample["managed_view_bytes"] = int(line.split("=")[1])
                        verify(paths[name], rows, expected)
                        samples[name].append(sample)
                # Native header read -> streaming export verifies the reader independently.
                copied = data / f"views-copy-{rows}.xlsx"
                subprocess.run(list(map(str, [target, "copy", paths["openpyxl"], copied, rows])), check=True, env=dict(os.environ, TMPDIR=directory), capture_output=True)
                verify(copied, rows, expected)
                assert not list(Path(directory).iterdir())
            case = {"rows": rows, "operation": mode, "samples": samples, "medians": {name: {key: statistics.median(sample[key] for sample in values) for key in ("seconds", "cpu_seconds", "peak_rss_kib", "sampled_temp_peak_bytes")} for name, values in samples.items()}}
            report["cases"].append(case)
            print(rows, mode, case["medians"], flush=True)
    args.output.write_text(json.dumps(report, indent=2) + "\n")

if __name__ == "__main__": main()
