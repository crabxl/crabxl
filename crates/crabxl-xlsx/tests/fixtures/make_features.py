"""Generate features.xlsx with openpyxl 3.1.5.

Run: uv run --with openpyxl==3.1.5 python make_features.py
"""
import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import (
    CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule, Rule,
)
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.packaging.custom import StringProperty, IntProperty, BoolProperty

wb = Workbook()
ws = wb.active
ws.title = "Data"
ws.append(["Name", "Qty", "Price"])
for i in range(1, 11):
    ws.append([f"item{i}", i, i * 1.5])
table = Table(displayName="Items", ref="A1:C11")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
ws.add_table(table)
ws.conditional_formatting.add("B2:B11", CellIsRule(operator="greaterThan", formula=["5"], fill=PatternFill("solid", bgColor="FFC7CE")))
ws.conditional_formatting.add("C2:C11", ColorScaleRule(start_type="min", start_color="FF0000", end_type="max", end_color="00FF00"))
ws.conditional_formatting.add("C2:C11", DataBarRule(start_type="min", end_type="max", color="638EC6"))
ws.conditional_formatting.add("B2:B11", IconSetRule("3Arrows", "percent", [0, 33, 67]))
ws.conditional_formatting.add("A2:A11", FormulaRule(formula=['LEN(A2)>5'], font=Font(bold=True)))
dxf = DifferentialStyle(font=Font(color="9C0006"))
ws.conditional_formatting.add("A2:A11", Rule(type="containsText", operator="containsText", text="item1", dxf=dxf, formula=['NOT(ISERROR(SEARCH("item1",A2)))']))
dv = DataValidation(type="list", formula1='"a,b,c"', allow_blank=True, showErrorMessage=True, error="Bad", errorTitle="Oops")
dv.add("E2:E20")
ws.add_data_validation(dv)
dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="10", prompt="1-10", promptTitle="Range")
dv2.add("F2:F20")
ws.add_data_validation(dv2)
ws["A1"].comment = Comment("Header note", "Alice")
c = Comment("Second line\nmore", "Bob", width=200, height=100)
ws["B3"].comment = c
ws.protection.sheet = True
ws.protection.password = "secret"
ws.protection.formatColumns = False
ws.oddHeader.center.text = "Report &[Page]"
ws.oddFooter.right.text = "Confidential"
ws.print_area = "A1:C11"
ws.print_title_rows = "1:1"

ws2 = wb.create_sheet("Filter")
ws2.append(["k", "v"])
for i in range(5):
    ws2.append([f"k{i}", i])
ws2.auto_filter.ref = "A1:B6"
ws2.auto_filter.add_filter_column(1, ["1", "2"])
ws2.auto_filter.add_sort_condition("B2:B6", descending=True)

wb.defined_names["Total"] = DefinedName("Total", attr_text="Data!$B$2:$B$11")
ws2.defined_names["LocalName"] = DefinedName("LocalName", attr_text="Filter!$A$1")
wb.security.workbookPassword = "pw"
wb.security.lockStructure = True
wb.properties.title = "Features"
wb.properties.creator = "CrabXL tests"
wb.properties.created = datetime.datetime(2026, 1, 2, 3, 4, 5)
wb.properties.modified = datetime.datetime(2026, 1, 2, 3, 4, 5)
wb.custom_doc_props.append(StringProperty(name="Project", value="crabxl"))
wb.custom_doc_props.append(IntProperty(name="Build", value=42))
wb.custom_doc_props.append(BoolProperty(name="Draft", value=True))
wb.save(Path(__file__).with_name("features.xlsx"))
