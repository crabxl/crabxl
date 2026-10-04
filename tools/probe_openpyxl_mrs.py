import io
import json
import warnings
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.datavalidation import DataValidation

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
XMLNS = "http://www.w3.org/XML/1998/namespace"


def save(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def sheet_xml(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return ET.fromstring(z.read("xl/worksheets/sheet1.xml"))


results = {"version": openpyxl.__version__}
wb = Workbook()
ws = wb.active
ws["A1"] = "=SUM(1,2)"
ws["A2"] = CellRichText(TextBlock(InlineFont(b=True), " "), "x", "")
ws.add_data_validation(DataValidation(type="whole"))
data = save(wb)
root = sheet_xml(data)
formula = root.find(f'.//{{{NS}}}c[@r="A1"]')
results["formula_has_empty_v"] = (
    formula.find(f"{{{NS}}}v") is not None and formula.find(f"{{{NS}}}v").text is None
)
rich = root.find(f'.//{{{NS}}}c[@r="A2"]')
results["rich_text_runs"] = [
    {"text": t.text, "xml_space": t.get(f"{{{XMLNS}}}space")}
    for t in rich.iter(f"{{{NS}}}t")
]
dv = root.find(f"{{{NS}}}dataValidations")
results["empty_validation_container"] = None if dv is None else dv.attrib

# Synthetic unknown extension: check read -> unrelated edit -> write preservation.
wb2 = Workbook()
wb2.active["A1"] = 1
base = save(wb2)
b = io.BytesIO()
with (
    zipfile.ZipFile(io.BytesIO(base)) as source,
    zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as target,
):
    for entry in source.infolist():
        payload = source.read(entry.filename)
        if entry.filename == "xl/worksheets/sheet1.xml":
            payload = payload.replace(
                b"</worksheet>",
                b'<extLst><ext uri="{OPENRSXL-PROBE}"><probe:marker xmlns:probe="urn:openrsxl:probe" value="keep"/></ext></extLst></worksheet>',
            )
        target.writestr(entry, payload)
with warnings.catch_warnings(record=True) as caught:
    warnings.simplefilter("always")
    loaded = load_workbook(io.BytesIO(b.getvalue()))
    loaded.active["A1"] = 2
    output = save(loaded)
results["unknown_extension_preserved"] = (
    sheet_xml(output).find(f"{{{NS}}}extLst") is not None
)
results["extension_warnings"] = [str(w.message) for w in caught]
try:
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image

    image = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(image, format="PNG")
    image.seek(0)
    wb3 = Workbook()
    wb3.active.add_image(XLImage(image))
    original = save(wb3)
    loaded = load_workbook(io.BytesIO(original))
    save(loaded)
    try:
        save(loaded)
        results["loaded_image_second_save"] = "success"
    except (ValueError, OSError) as e:
        results["loaded_image_second_save"] = f"{type(e).__name__}: {e}"
except ImportError:
    results["loaded_image_second_save"] = "not tested: Pillow unavailable"
print(json.dumps(results, ensure_ascii=False, indent=2))
